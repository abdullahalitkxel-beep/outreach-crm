#!/usr/bin/env python3
"""
Outreach bot: personalized cold-email sequences with follow-ups, reply detection,
bounce handling and a suppression list.

Commands:
  python bot.py import leads.csv      # add leads
  python bot.py run [--dry-run]       # check replies, then send whatever is due
  python bot.py check                 # only check inboxes for replies/bounces
  python bot.py stats                 # campaign numbers
  python bot.py suppress a@b.com      # never email this address
"""
import argparse
import csv
import email
import imaplib
import io
import json
import logging
import os
import random
import re
import smtplib
import socket
import sqlite3
import sys
import time
from datetime import date, datetime, timedelta, timezone
from email.message import EmailMessage
from email.utils import formatdate, make_msgid, parseaddr
from pathlib import Path
from zoneinfo import ZoneInfo

import requests
import yaml
from dotenv import load_dotenv

# ROOT = where the code (and default config.yaml) lives.
# DATA = where YOUR data lives: database, settings, keys. In the Windows/Mac app this is
#        %APPDATA%\OutreachCRM (or ~/Library/Application Support/OutreachCRM); from source it's this folder.
ROOT = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))


def _package_family():
    """Package family name when running as a Microsoft Store (MSIX) app, else None."""
    if sys.platform != "win32":
        return None
    try:
        import ctypes
        from ctypes import wintypes
        n = wintypes.UINT(0)
        if ctypes.windll.kernel32.GetCurrentPackageFamilyName(ctypes.byref(n), None) == 15700:  # no package
            return None
        buf = ctypes.create_unicode_buffer(n.value)
        if ctypes.windll.kernel32.GetCurrentPackageFamilyName(ctypes.byref(n), buf) == 0:
            return buf.value
    except Exception:
        pass
    return None


def _data_dir():
    pfn = None if os.getenv("OUTREACH_HOME") else _package_family()
    if os.getenv("OUTREACH_HOME"):
        d = Path(os.environ["OUTREACH_HOME"])
    elif pfn:  # Store app: the package's own data folder (removed cleanly on uninstall)
        d = Path(os.getenv("LOCALAPPDATA", Path.home())) / "Packages" / pfn / "LocalState" / "OutreachCRM"
    elif getattr(sys, "frozen", False):
        if sys.platform == "win32":
            d = Path(os.getenv("APPDATA", Path.home())) / "OutreachCRM"
        elif sys.platform == "darwin":
            d = Path.home() / "Library" / "Application Support" / "OutreachCRM"
        else:
            d = Path.home() / ".outreachcrm"
    else:
        d = Path(__file__).resolve().parent
    d.mkdir(parents=True, exist_ok=True)
    for name in ("config.yaml", ".env.example"):  # first run of the app: copy defaults
        if not (d / name).exists() and (ROOT / name).exists():
            (d / name).write_bytes((ROOT / name).read_bytes())
    return d


DATA = _data_dir()
load_dotenv(DATA / ".env")  # non-secret overrides only; secrets live encrypted in secrets.dat

from secretstore import SecretStore  # noqa: E402

SECRETS = SecretStore(DATA)
try:
    SECRETS.migrate_env_file(DATA / ".env")
except Exception as _e:
    print("secrets migration:", _e, file=sys.stderr)
SECRETS.apply_to_environ()
log = logging.getLogger("outreach")

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
FIND_EMAILS_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
CORE_FIELDS = {"email", "first_name", "company", "website"}
REPLY_LABELS = ["interested", "meeting_request", "question", "not_interested", "unsubscribe", "out_of_office", "referral", "other"]
HANDOVER_LABELS = ("interested", "meeting_request", "question")


# ---------------------------------------------------------------- helpers

def env(name):
    """Read a secret from .env, treating untouched placeholders (sk-ant-..., your-app-password) as missing."""
    v = (os.getenv(name) or "").strip()
    return "" if (not v or "..." in v or v.lower().startswith("your")) else v


def now_utc():
    return datetime.now(timezone.utc)


def iso(dt):
    return dt.astimezone(timezone.utc).isoformat(timespec="seconds")


class SafeDict(dict):
    """format_map dict that leaves unknown placeholders blank instead of crashing."""
    def __missing__(self, key):
        return ""


def _deep_merge(base, over):
    for k, v in (over or {}).items():
        if isinstance(v, dict) and isinstance(base.get(k), dict):
            _deep_merge(base[k], v)
        else:
            base[k] = v
    return base


def settings_path(config_path):
    return Path(config_path).with_name("settings.yaml")


def load_config(path):
    """config.yaml (defaults, hand-editable) overlaid with settings.yaml (written by the app's Settings page)."""
    with open(path, encoding="utf-8") as f:
        cfg = yaml.safe_load(f)
    sp = settings_path(path)
    if sp.exists():
        with open(sp, encoding="utf-8") as f:
            _deep_merge(cfg, yaml.safe_load(f) or {})
    return cfg


# ---------------------------------------------------------------- database

SCHEMA = """
CREATE TABLE IF NOT EXISTS leads (
    id INTEGER PRIMARY KEY,
    email TEXT UNIQUE NOT NULL,
    first_name TEXT, company TEXT, website TEXT,
    extra TEXT DEFAULT '{}',
    status TEXT DEFAULT 'active',      -- active | done | replied | not_interested | unsubscribed | bounced
    step INTEGER DEFAULT 0,            -- index of the next sequence step to send
    next_send_at TEXT,
    sender TEXT,                       -- inbox that owns this lead (follow-ups use the same one)
    personal_line TEXT,
    first_msg_id TEXT,
    first_subject TEXT,
    created_at TEXT,
    stage TEXT DEFAULT 'new',          -- CRM pipeline stage (see STAGES)
    deal_value REAL DEFAULT 0
);
CREATE TABLE IF NOT EXISTS notes (
    id INTEGER PRIMARY KEY, lead_id INTEGER, body TEXT, at TEXT
);
CREATE TABLE IF NOT EXISTS events (
    id INTEGER PRIMARY KEY, lead_id INTEGER, sender TEXT, type TEXT, detail TEXT, at TEXT
);
CREATE TABLE IF NOT EXISTS suppression (email TEXT PRIMARY KEY, reason TEXT, at TEXT);
CREATE TABLE IF NOT EXISTS seen_messages (msg_id TEXT PRIMARY KEY);
CREATE TABLE IF NOT EXISTS outbox (                -- drafts waiting for review (review mode)
    id INTEGER PRIMARY KEY, lead_id INTEGER, step INTEGER, sender TEXT,
    subject TEXT, body TEXT,
    status TEXT,                                   -- pending | approved | sent | rejected | cancelled
    created_at TEXT, decided_at TEXT, sent_at TEXT
);
CREATE TABLE IF NOT EXISTS accounts (
    id INTEGER PRIMARY KEY, name TEXT, domain TEXT, industry TEXT, country TEXT, created_at TEXT
);
CREATE TABLE IF NOT EXISTS opportunities (
    id INTEGER PRIMARY KEY, name TEXT, account_id INTEGER, lead_id INTEGER,
    stage TEXT, amount REAL DEFAULT 0, close_date TEXT, next_step TEXT,
    created_at TEXT, closed_at TEXT
);
CREATE TABLE IF NOT EXISTS tasks (
    id INTEGER PRIMARY KEY, lead_id INTEGER, opp_id INTEGER, subject TEXT,
    due_date TEXT, status TEXT DEFAULT 'open', created_at TEXT, done_at TEXT
);
CREATE TABLE IF NOT EXISTS jobs (
    id INTEGER PRIMARY KEY, cmd TEXT, args TEXT, status TEXT, output TEXT, started_at TEXT, finished_at TEXT
);
CREATE TABLE IF NOT EXISTS messages (                -- a supervisor's emails with handed-over prospects
    id INTEGER PRIMARY KEY, lead_id INTEGER, msg_id TEXT UNIQUE, direction TEXT, at TEXT, frm TEXT, to_ TEXT, cc TEXT,
    subject TEXT, body TEXT, seen INTEGER DEFAULT 0, via TEXT
);
CREATE TABLE IF NOT EXISTS campaigns (                   -- separate offers / segments, each with its own sequence and inboxes
    id INTEGER PRIMARY KEY, name TEXT, status TEXT DEFAULT 'active', inboxes TEXT DEFAULT '[]', sequence TEXT,
    offer TEXT, created_at TEXT
);
CREATE TABLE IF NOT EXISTS inbox_health (                -- inboxes the deliverability guard has paused
    email TEXT PRIMARY KEY, paused INTEGER DEFAULT 0, reason TEXT, at TEXT, login_failures INTEGER DEFAULT 0
);
CREATE TABLE IF NOT EXISTS comments (                    -- comments on prospects between team members (with @mentions)
    id INTEGER PRIMARY KEY, uid TEXT UNIQUE, lead_email TEXT, lead_name TEXT, company TEXT, author TEXT, author_name TEXT,
    text TEXT, at TEXT, to_ TEXT, mine INTEGER DEFAULT 0, synced INTEGER DEFAULT 0, seen INTEGER DEFAULT 0
);
CREATE TABLE IF NOT EXISTS domain_health (               -- SPF / DKIM / DMARC / MX checks of the sending domains
    domain TEXT PRIMARY KEY, at TEXT, result TEXT
);
CREATE TABLE IF NOT EXISTS instantly_pulls (            -- SuperSearch imports started by the bot
    id INTEGER PRIMARY KEY, resource_id TEXT, started_at TEXT, finished_at TEXT, wanted INTEGER, status TEXT,
    found INTEGER, imported INTEGER, skipped TEXT
);
CREATE TABLE IF NOT EXISTS discovered (                   -- companies already searched for published contacts
    domain TEXT PRIMARY KEY, at TEXT, found INTEGER, added INTEGER, note TEXT
);
CREATE TABLE IF NOT EXISTS handover_out (                -- handovers waiting to reach the team hub
    id INTEGER PRIMARY KEY, lead_id INTEGER, supervisor TEXT, payload TEXT, at TEXT, synced INTEGER DEFAULT 0
);
CREATE TABLE IF NOT EXISTS apollo_seen (apollo_id TEXT PRIMARY KEY, result TEXT, at TEXT);
CREATE TABLE IF NOT EXISTS apollo_progress (search TEXT PRIMARY KEY, page INTEGER);
"""


# Lead path (Salesforce-style): open stages, then closed as converted or unqualified
LEAD_OPEN_STAGES = ["new", "contacted", "replied", "qualified"]
STAGES = LEAD_OPEN_STAGES + ["converted", "unqualified"]
EARLY_STAGES = ("new", "contacted")  # the bot only moves leads that a human hasn't advanced
HUMAN_STAGES = ("qualified", "converted", "unqualified")  # a person has taken over: stop emails

# Opportunity path with default win probability (%)
OPP_STAGES = {"Discovery": 20, "Proposal": 50, "Negotiation": 75, "Closed Won": 100, "Closed Lost": 0}
OPP_OPEN_STAGES = ["Discovery", "Proposal", "Negotiation"]


def _add_column(con, table, col, decl):
    cols = {r["name"] for r in con.execute(f"PRAGMA table_info({table})")}
    if col not in cols:
        con.execute(f"ALTER TABLE {table} ADD COLUMN {col} {decl}")


_migrated = set()


def db_connect(path):
    """Open the database. WAL mode lets the CRM and bot runs read and write at the same time;
    the schema check/upgrade runs only once per process."""
    full = DATA / path
    con = sqlite3.connect(full, check_same_thread=False, timeout=30)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA busy_timeout = 30000")
    if str(full) not in _migrated:
        con.execute("PRAGMA journal_mode = WAL")
        con.executescript(SCHEMA)
        # upgrade databases created by earlier versions
        _add_column(con, "leads", "stage", "TEXT DEFAULT 'new'")
        _add_column(con, "leads", "deal_value", "REAL DEFAULT 0")
        _add_column(con, "leads", "account_id", "INTEGER")
        _add_column(con, "notes", "kind", "TEXT DEFAULT 'note'")
        _add_column(con, "outbox", "variant", "TEXT")
        _add_column(con, "events", "variant", "TEXT")
        _add_column(con, "leads", "handed_to", "TEXT")          # supervisor this prospect was handed to (sender side)
        _add_column(con, "leads", "handed_at", "TEXT")
        _add_column(con, "leads", "handed_from", "TEXT")        # who handed it over (supervisor side)
        _add_column(con, "leads", "prospect_stage", "TEXT")     # supervisor's stage for a handed-over prospect
        _add_column(con, "leads", "campaign_id", "INTEGER DEFAULT 1")
        _add_column(con, "inbox_health", "resumed_at", "TEXT")
        _add_column(con, "instantly_pulls", "campaign_id", "INTEGER DEFAULT 1")
        _add_column(con, "campaigns", "brief", "TEXT")
        _add_column(con, "campaigns", "writing", "TEXT")
        con.execute("INSERT OR IGNORE INTO campaigns(id, name, status, created_at) VALUES (1, 'Main campaign', 'active', ?)", (iso(now_utc()),))
        con.execute("UPDATE leads SET stage='qualified' WHERE stage IN ('meeting','proposal','won')")
        con.execute("UPDATE leads SET stage='unqualified' WHERE stage='lost'")
        import convo
        convo.ensure(con, _add_column)
        import report
        report.ensure(con, _add_column)
        import autopilot
        autopilot.ensure(con, _add_column)
        con.execute("CREATE INDEX IF NOT EXISTS sent_log_email ON sent_log(lead_email)")
        con.execute("CREATE INDEX IF NOT EXISTS events_type_sender_at ON events(type, sender, at)")     # hundreds of inboxes: fast daily counts
        con.execute("CREATE INDEX IF NOT EXISTS sent_log_domain_at ON sent_log(domain, at)")
        con.execute("CREATE INDEX IF NOT EXISTS leads_due ON leads(status, next_send_at)")
        _test_triggers(con)
        con.commit()
        _migrated.add(str(full))
    return con


# ---------------------------------------------------------------- test mode kept apart from live
# While test mode is on, every change to a lead is recorded: the lead as it was before the first change, and every
# event, draft, note, task, handover and do-not-contact entry added. Going live (or "Clear test results") puts the
# leads back exactly as they were and removes what the test added, so test replies never close a real lead.

TEST_TRACKED = ("events", "outbox", "notes", "tasks", "handover_out", "suppression", "sent_log", "received", "temp_log",
                "related_contacts", "lead_addresses")
# what a test run changes on a lead; edits you make yourself (name, company, notes on the lead itself) are kept
TEST_PROGRESS = {"step", "status", "stage", "next_send_at", "sender", "first_msg_id", "first_subject", "handed_to", "handed_at", "prospect_stage",
                 "temperature", "temp_reason", "check_back", "came_back", "score", "deal_status", "lost_reason", "journey"}


def _test_triggers(con):
    con.executescript("""
CREATE TABLE IF NOT EXISTS app_mode (k TEXT PRIMARY KEY, v TEXT);
CREATE TABLE IF NOT EXISTS test_snapshot (lead_id INTEGER PRIMARY KEY, data TEXT, at TEXT);
CREATE TABLE IF NOT EXISTS test_rows (tbl TEXT, rid TEXT, PRIMARY KEY (tbl, rid));
""")
    cols = [r[1] for r in con.execute("PRAGMA table_info(leads)") if r[1] != "id"]
    snap = "json_object(" + ", ".join(f"'{c}', OLD.{c}" for c in cols) + ")"
    on = "(SELECT v FROM app_mode WHERE k='test')='1'"
    con.execute("DROP TRIGGER IF EXISTS test_snap_leads")
    con.execute(f"CREATE TRIGGER test_snap_leads BEFORE UPDATE ON leads WHEN {on} BEGIN "
                f"INSERT OR IGNORE INTO test_snapshot(lead_id, data, at) VALUES (OLD.id, {snap}, strftime('%Y-%m-%dT%H:%M:%SZ','now')); END")
    for t in TEST_TRACKED:
        key = "NEW.email" if t == "suppression" else "NEW.rowid"
        con.execute(f"DROP TRIGGER IF EXISTS test_rows_{t}")
        con.execute(f"CREATE TRIGGER test_rows_{t} AFTER INSERT ON {t} WHEN {on} BEGIN "
                    f"INSERT OR IGNORE INTO test_rows(tbl, rid) VALUES ('{t}', {key}); END")


def test_on(cfg):
    tm = cfg.get("test_mode") or {}
    return bool(tm.get("enabled"))


def sync_mode(con, cfg):
    """Tell the database which mode we're in. Switching from test to live clears the test results."""
    want = "1" if test_on(cfg) else "0"
    row = con.execute("SELECT v FROM app_mode WHERE k='test'").fetchone()
    if row and row[0] == want:
        return
    if row and row[0] == "1" and want == "0":
        clear_test_results(con)
    if not row:
        _adopt_old_tests(con)
    con.execute("INSERT INTO app_mode(k, v) VALUES ('test', ?) ON CONFLICT(k) DO UPDATE SET v=excluded.v", (want,))
    con.commit()


def _adopt_old_tests(con):
    """Tests sent before this version weren't recorded. Leads whose only emails so far were test emails are
    marked to start fresh at go-live, and those test emails are removed with the other test results."""
    for l in con.execute("SELECT * FROM leads WHERE id IN (SELECT lead_id FROM events WHERE type='sent' AND detail LIKE '%(test)%') "
                         "AND id NOT IN (SELECT lead_id FROM events WHERE type='sent' AND detail NOT LIKE '%(test)%')").fetchall():
        fresh = {"step": 0, "status": "active" if l["status"] in ("active", "done", "paused") else l["status"], "next_send_at": None, "sender": None,
                 "first_msg_id": None, "first_subject": None, "stage": "new" if l["stage"] == "contacted" else l["stage"]}
        con.execute("INSERT OR IGNORE INTO test_snapshot(lead_id, data, at) VALUES (?,?,?)", (l["id"], json.dumps(fresh), iso(now_utc())))
    con.execute("INSERT OR IGNORE INTO test_rows(tbl, rid) SELECT 'events', CAST(id AS TEXT) FROM events WHERE type='sent' AND detail LIKE '%(test)%'")


def test_counts(con):
    return {"leads": con.execute("SELECT COUNT(*) FROM test_snapshot").fetchone()[0],
            "events": con.execute("SELECT COUNT(*) FROM test_rows WHERE tbl='events'").fetchone()[0]}


def clear_test_results(con):
    """Put every lead back as it was before test mode touched it and remove what the test added. Returns leads restored."""
    was = con.execute("SELECT v FROM app_mode WHERE k='test'").fetchone()
    con.execute("UPDATE app_mode SET v='0' WHERE k='test'")      # so the restore itself isn't recorded
    cols = {r[1] for r in con.execute("PRAGMA table_info(leads)")}
    n = 0
    for r in con.execute("SELECT lead_id, data FROM test_snapshot").fetchall():
        data = {k: v for k, v in json.loads(r["data"] or "{}").items() if k in cols and k in TEST_PROGRESS}
        if data:
            con.execute(f"UPDATE leads SET {', '.join(k + '=?' for k in data)} WHERE id=?", (*data.values(), r["lead_id"]))
            n += 1
    for t in TEST_TRACKED:
        if t == "suppression":
            con.execute("DELETE FROM suppression WHERE email IN (SELECT rid FROM test_rows WHERE tbl='suppression')")
        else:
            con.execute(f"DELETE FROM {t} WHERE rowid IN (SELECT CAST(rid AS INTEGER) FROM test_rows WHERE tbl=?)", (t,))
    con.execute("DELETE FROM test_snapshot")
    con.execute("DELETE FROM test_rows")
    if was:
        con.execute("UPDATE app_mode SET v=? WHERE k='test'", (was[0],))
    con.commit()
    return n


def is_test_row(con, tbl, rid):
    return con.execute("SELECT 1 FROM test_rows WHERE tbl=? AND rid=?", (tbl, str(rid))).fetchone() is not None


def add_task(con, subject, due_date=None, lead_id=None, opp_id=None):
    due = due_date or datetime.now().date().isoformat()
    con.execute("INSERT INTO tasks(lead_id, opp_id, subject, due_date, status, created_at) "
                "VALUES (?,?,?,?,'open',?)", (lead_id, opp_id, subject, due, iso(now_utc())))


def set_stage(con, lead_id, stage, auto=True):
    """Move a lead in the pipeline. Automatic moves never override a human's later stage."""
    if auto:
        cur = con.execute(
            f"UPDATE leads SET stage=? WHERE id=? AND stage IN ({','.join('?' * len(EARLY_STAGES))})",
            (stage, lead_id, *EARLY_STAGES),
        )
    else:
        cur = con.execute("UPDATE leads SET stage=? WHERE id=?", (stage, lead_id))
    if cur.rowcount:
        log_event(con, lead_id, "stage", stage)


def log_event(con, lead_id, type_, detail="", sender=None, variant=None):
    con.execute(
        "INSERT INTO events(lead_id, sender, type, detail, at, variant) VALUES (?,?,?,?,?,?)",
        (lead_id, sender, type_, detail, iso(now_utc()), variant),
    )


def cancel_drafts(con, lead_id):
    """Withdraw any unsent drafts for a lead (they replied, unsubscribed, were paused...)."""
    con.execute("UPDATE outbox SET status='cancelled', decided_at=? "
                "WHERE lead_id=? AND status IN ('pending','approved')", (iso(now_utc()), lead_id))


def suppress(con, addr, reason):
    con.execute(
        "INSERT OR IGNORE INTO suppression(email, reason, at) VALUES (?,?,?)",
        (addr.lower(), reason, iso(now_utc())),
    )


# ---------------------------------------------------------------- senders

def ramp_limit(cfg, s, today=None):
    """New inboxes start low and climb each week up to their daily limit (protects a young domain's reputation)."""
    full = int(s.get("daily_limit") or 10)
    if not s.get("ramp_from"):
        return full
    rc = cfg.get("ramp") or {}
    try:
        start = datetime.fromisoformat(str(s["ramp_from"])[:10]).date()
    except ValueError:
        return full
    weeks = max(0, ((today or datetime.now().date()) - start).days // 7)
    return min(full, int(rc.get("start", 5)) + int(rc.get("per_week", 5)) * weeks)


def load_senders(cfg):
    senders = []
    for s in cfg["senders"]:
        pw = env(s["password_env"])
        if not pw:
            log.warning("No password in env var %s for %s", s["password_env"], s["email"])
        senders.append({**s, "password": pw, "full_limit": int(s.get("daily_limit") or 10), "daily_limit": ramp_limit(cfg, s)})
    return senders


# ---------------------------------------------------------------- deliverability guard

def guard_paused(con):
    return {r["email"].lower(): r["reason"] for r in con.execute("SELECT email, reason FROM inbox_health WHERE paused=1")}


def guard_pause(cfg, con, email, reason):
    con.execute("INSERT INTO inbox_health(email, paused, reason, at) VALUES (?,1,?,?) "
                "ON CONFLICT(email) DO UPDATE SET paused=1, reason=excluded.reason, at=excluded.at", (email.lower(), reason, iso(now_utc())))
    con.commit()
    log.error("Inbox %s paused: %s", email, reason)
    notify(f"⛔ BSO paused the inbox {email}: {reason} It sends nothing until you resume it in Settings → Email inboxes.", cfg)


def guard_check(cfg, con, senders):
    """Pause inboxes whose recent bounce rate is too high. Returns the problems found."""
    g = cfg.get("guard") or {}
    if g.get("enabled", True) is False:
        return []
    limit, min_sent = float(g.get("bounce_rate", 0.05)), int(g.get("min_sent", 20))
    paused, out = guard_paused(con), []
    since_default = iso(now_utc() - timedelta(days=7))
    for s in senders:
        e = s["email"].lower()
        if e in paused:
            continue
        r = con.execute("SELECT resumed_at FROM inbox_health WHERE email=?", (e,)).fetchone()
        since = max(since_default, r["resumed_at"]) if r and r["resumed_at"] else since_default
        op = ">" if r and r["resumed_at"] and r["resumed_at"] >= since_default else ">="     # after a resume: only what happened since
        sent = con.execute(f"SELECT COUNT(*) FROM events WHERE type='sent' AND lower(sender)=? AND at{op}? AND detail NOT LIKE '%(test)%'", (e, since)).fetchone()[0]
        bounced = con.execute(f"SELECT COUNT(*) FROM events WHERE type='bounce' AND lower(sender)=? AND at{op}?", (e, since)).fetchone()[0]
        if sent >= min_sent and bounced / sent >= limit:
            reason = f"{bounced} of its last {sent} emails bounced ({bounced * 100 // sent}%, the limit is {limit * 100:g}%)."
            guard_pause(cfg, con, e, reason)
            out.append(f"{s['email']}: paused, {reason}")
    return out


def guard_login(cfg, con, email, ok):
    """Count login failures; pause after 3 in a row so a broken inbox isn't retried forever."""
    e = email.lower()
    if ok:
        con.execute("UPDATE inbox_health SET login_failures=0 WHERE email=?", (e,))
        return
    con.execute("INSERT INTO inbox_health(email, login_failures) VALUES (?,1) ON CONFLICT(email) DO UPDATE SET login_failures=login_failures+1", (e,))
    n = con.execute("SELECT login_failures FROM inbox_health WHERE email=?", (e,)).fetchone()[0]
    if n >= int((cfg.get("guard") or {}).get("login_failures", 3)):
        guard_pause(cfg, con, e, f"its login was rejected {n} times in a row (password changed or mailbox locked?).")


def sent_today(con, sender_email, cfg):
    tz = ZoneInfo(cfg["schedule"]["timezone"])
    start = datetime.now(tz).replace(hour=0, minute=0, second=0, microsecond=0)
    row = con.execute(
        "SELECT COUNT(*) FROM events WHERE type='sent' AND sender=? AND at>=?",
        (sender_email, iso(start)),
    ).fetchone()
    return row[0]


def in_send_window(cfg):
    w = cfg["schedule"]
    t = datetime.now(ZoneInfo(w["timezone"]))
    return t.isoweekday() in w["days"] and w["start_hour"] <= t.hour < w["end_hour"]


# ---------------------------------------------------------------- LLM

def _reply_text(msg):
    """The answer text of a Messages API reply (newer models put thinking blocks before it)."""
    return "".join(getattr(b, "text", "") for b in (msg.content or []) if getattr(b, "type", "") == "text").strip()


AI_RATES = {"haiku": (1.0, 5.0), "sonnet": (3.0, 15.0), "opus": (5.0, 25.0)}   # approximate USD per million tokens; editable in Settings


def ai_spend_path():
    return DATA / "ai_spend.json"


def ai_spend(month=None):
    try:
        d = json.loads(ai_spend_path().read_text(encoding="utf-8"))
    except Exception:
        d = {}
    return d.get(month or date.today().strftime("%Y-%m"), {"usd": 0.0, "in": 0, "out": 0, "calls": 0})


def _record_spend(cfg, model, usage):
    try:
        rates = dict(AI_RATES, **((cfg.get("ai_budget") or {}).get("rates") or {}))
        rin, rout = next((v for k, v in rates.items() if k in (model or "").lower()), (3.0, 15.0))
        tin, tout = int(getattr(usage, "input_tokens", 0) or 0), int(getattr(usage, "output_tokens", 0) or 0)
        p = ai_spend_path()
        try:
            d = json.loads(p.read_text(encoding="utf-8"))
        except Exception:
            d = {}
        m = date.today().strftime("%Y-%m")
        cur = d.get(m) or {"usd": 0.0, "in": 0, "out": 0, "calls": 0}
        cur.update(usd=round(cur["usd"] + tin / 1e6 * float(rin) + tout / 1e6 * float(rout), 4), calls=cur["calls"] + 1,
                   **{"in": cur["in"] + tin, "out": cur["out"] + tout})
        d[m] = cur
        tmp = p.with_suffix(".tmp")
        tmp.write_text(json.dumps(d), encoding="utf-8")
        os.replace(tmp, p)
    except Exception as e:
        log.warning("AI spend not recorded: %s", e)


def ai_over_budget(cfg):
    cap = float((cfg.get("ai_budget") or {}).get("monthly_usd") or 0)
    return bool(cap) and ai_spend()["usd"] >= cap


def llm(cfg, model_key, prompt, max_tokens=200):
    """Call Claude; returns text or None if unavailable/failing (or over this month's AI budget)."""
    if not env("ANTHROPIC_API_KEY"):
        return None
    if ai_over_budget(cfg):
        log.warning("Monthly AI budget reached: using templates and keyword rules until next month (Settings → AI & Apollo keys).")
        return None
    try:
        import anthropic
        client = anthropic.Anthropic(api_key=env("ANTHROPIC_API_KEY"))
        model = (cfg.get("models") or {}).get(model_key) or (cfg.get("models") or {}).get("personalize")
        try:
            # short, simple writing: no thinking needed (Haiku accepts this; others may refuse it)
            msg = client.messages.create(model=model, max_tokens=max_tokens, messages=[{"role": "user", "content": prompt}],
                                         extra_body={"thinking": {"type": "disabled"}})
        except anthropic.BadRequestError:
            # model can't switch thinking off: leave room for it, keep only the answer
            msg = client.messages.create(model=model, max_tokens=max_tokens + 4000, messages=[{"role": "user", "content": prompt}])
        _record_spend(cfg, model, getattr(msg, "usage", None))
        text = _reply_text(msg)
        return text or None
    except Exception as e:  # network, quota, etc. -> fall back gracefully
        log.warning("LLM call failed: %s", e)
        return None


def fetch_site_text(url, limit=4000):
    if not url:
        return ""
    if not url.startswith("http"):
        url = "https://" + url
    try:
        from bs4 import BeautifulSoup
        r = requests.get(url, timeout=10, headers={"User-Agent": "Mozilla/5.0"})
        soup = BeautifulSoup(r.text, "html.parser")
        for tag in soup(["script", "style", "noscript", "svg"]):
            tag.decompose()
        return " ".join(soup.get_text(" ").split())[:limit]
    except Exception as e:
        log.info("Could not fetch %s: %s", url, e)
        return ""



_AVOID_CACHE = {}


def _avoid_patterns(cfg):
    phrases = tuple(cfg.get("avoid_phrases") or ())
    if phrases not in _AVOID_CACHE:
        pats = []
        for ph in phrases:
            ph = str(ph).strip().replace("\u2019", "'")
            if not ph:
                continue
            if ph.startswith("regex:"):
                pats.append((ph, re.compile(ph[6:], re.I)))
            elif re.search(r"[$%]", ph):          # symbols: match literally
                pats.append((ph, re.compile(re.escape(ph), re.I)))
            else:                                  # whole words; spaces and hyphens are interchangeable
                words = [re.escape(w) for w in re.split(r"[\s-]+", ph) if w]
                pats.append((ph, re.compile(r"(?<![\w])" + r"[\s-]+".join(words) + r"(?![\w])", re.I)))
        _AVOID_CACHE[phrases] = pats
    return _AVOID_CACHE[phrases]


def find_avoid(cfg, *texts):
    """Phrases from the avoid list (config: avoid_phrases) that appear in the text, e.g. spam-trigger words."""
    text = "\n".join(t or "" for t in texts).replace("\u2019", "'")
    return sorted({ph.replace("regex:", "") for ph, rx in _avoid_patterns(cfg) if rx.search(text)}, key=str.lower)


def template_avoid_hits(cfg):
    """Avoid-list phrases found in the email templates after the facts are filled in. {step: {variant: [phrases]}}"""
    out = {}
    facts = SafeDict(**{k: str(v) for k, v in (cfg.get("facts") or {}).items()})
    for i, st in enumerate(cfg.get("sequence") or [], 1):
        for v in st.get("variants") or [st]:
            if v.get("enabled", True) is False:
                continue
            filled = ((v.get("subject") or "") + "\n" + (v.get("body") or "")).format_map(facts)
            hits = find_avoid(cfg, filled, cfg.get("footer", ""))
            if hits:
                out.setdefault(i, {})[v.get("name") or f"step {i}"] = hits
    return out


def personalize(cfg, lead):
    """One personalised opening line. Privacy: the AI only sees business information — the company's
    public website text, company name, job title and industry — never the contact's name, email,
    LinkedIn or other personal details."""
    pc = cfg["personalization"]
    fallback = pc["fallback_line"].format_map(SafeDict(company=short_company(lead["company"]) or "your team"))
    if not pc.get("enabled"):
        return fallback
    site = fetch_site_text(lead["website"])
    extra = json.loads(lead["extra"] or "{}")
    business = {k: extra[k] for k in ("title", "industry") if extra.get(k)}
    prompt = f"""You write the opening line of a B2B cold email.

What we offer: {cfg['offer']}

The prospect's company:
- Company: {short_company(lead['company']) or 'unknown'}
- Recipient's job title and industry: {json.dumps(business) if business else 'unknown'}
- Text from the company's public website: {site or 'not available'}

Write ONE sentence (max 30 words) addressed to the recipient that shows we looked at their company
specifically and leads naturally toward our offer. Rules: reference a concrete detail from the info
above; don't use any person's name; no flattery clichés; no questions; never invent facts.
If there isn't enough specific information, output exactly: NONE
Output only the sentence."""
    line = llm(cfg, "personalize", prompt)
    if not line or line.upper().startswith("NONE") or not (10 < len(line) < 300):
        return fallback
    line = line.strip().strip('"')
    if find_avoid(cfg, line):  # the AI used a word you asked us to avoid
        return fallback
    return line


def analyze_reply(cfg, msg, body):
    """{'label', 'temperature', 'check_back', 'reason'} from the AI when it's on (else just the label from keywords)."""
    low = body.lower()
    if any(k in low for k in ("unsubscribe", "remove me", "stop emailing", "take me off")):
        return {"label": "unsubscribe", "temperature": "closed"}
    auto = (msg.get("Auto-Submitted", "") + msg.get("X-Autoreply", "")).lower()
    if auto and auto != "no":
        return {"label": "out_of_office"}
    if not (cfg.get("ai_reply_sorting") or {}).get("enabled"):
        return {"label": keyword_label(low)}
    prompt = f"""You sort replies to B2B cold sales emails. Today is {date.today().isoformat()}.
Return only JSON: {{"label": one of {REPLY_LABELS}, "temperature": "hot"|"warm"|"cold"|"closed",
"check_back": "YYYY-MM-DD" or null, "reason": "max 12 words"}}
label: interested = positive, wants to know more; meeting_request = asks for or proposes a call/meeting/time;
question = asks something specific (price, process, examples); referral = points us to someone else; other = needs a human look.
temperature: hot = ready to talk/buy now (meeting, price, proposal, timeline); warm = interested or asking; cold = not now, has a vendor,
no budget, polite no; closed = wants no more contact. check_back = the date they suggest coming back (e.g. "after Q1"), else null.

Reply:
\"\"\"{body[:2500]}\"\"\""""
    raw = llm(cfg, "classify", prompt, max_tokens=120)
    if raw:
        m = re.search(r"\{.*\}", raw, re.S)
        try:
            j = json.loads(m.group(0)) if m else {}
        except ValueError:
            j = {}
        lab = str(j.get("label") or "").lower().strip()
        if lab in REPLY_LABELS:
            return {"label": lab, "temperature": str(j.get("temperature") or "").lower(), "check_back": j.get("check_back"),
                    "reason": str(j.get("reason") or "")[:120]}
        lab = raw.lower().strip(" .\n")
        if lab in REPLY_LABELS:
            return {"label": lab}
    return {"label": keyword_label(low)}


def classify_reply(cfg, msg, body):
    low = body.lower()
    if any(k in low for k in ("unsubscribe", "remove me", "stop emailing", "take me off")):
        return "unsubscribe"
    auto = (msg.get("Auto-Submitted", "") + msg.get("X-Autoreply", "")).lower()
    if auto and auto != "no":
        return "out_of_office"
    if not (cfg.get("ai_reply_sorting") or {}).get("enabled"):
        return keyword_label(low)  # default: sorted on this computer; reply text isn't sent anywhere
    prompt = f"""Classify this reply to a cold sales email. Answer with exactly one label from:
{', '.join(REPLY_LABELS)}
(interested = positive, wants to know more; meeting_request = asks for or proposes a call/meeting/time;
question = asks something specific such as price, process or examples; referral = they point us to someone else;
other = anything that needs a human look)

Reply:
\"\"\"{body[:2000]}\"\"\""""
    label = llm(cfg, "classify", prompt, max_tokens=12)
    if label is None:
        return keyword_label(low)
    label = label.lower().strip(" .\n")
    return label if label in REPLY_LABELS else "other"


def keyword_label(low):
    """Fallback reply sorting when no AI key is set (or the AI is unavailable)."""
    if re.search(r"\b(out of (the )?office|on (annual )?leave|on vacation|on holiday|away until|limited access to (my )?email)\b", low):
        return "out_of_office"
    if re.search(r"\b(not interested|no thanks|no thank you|not a (good )?fit|we('re| are) (all )?set|not (looking|right now)|please don'?t (contact|email))\b", low):
        return "not_interested"
    if re.search(r"\b(reach out to|contact|speak (to|with)|talk to) [a-z]+ (instead|who|on our)|\bright person\b|\bnot the right (person|contact)\b", low):
        return "referral"
    if re.search(r"\b(schedule|book a|set up a (call|meeting|time)|(a|quick|short|intro) (call|chat|meeting)|meeting|calendar|calendly|zoom|teams call|google meet|"
                 r"what time|which time|(are|is) you (free|available)|i'?m (free|available)|available (on|next|this|tomorrow)|"
                 r"(monday|tuesday|wednesday|thursday|friday)( at| morning| afternoon)?|tomorrow at|next week)\b", low):
        return "meeting_request"
    if re.search(r"\b(interested|let'?s (talk|chat|connect|meet)|(can|could|shall) we (talk|chat|meet|speak|connect)|happy to (talk|chat|connect)|sounds (good|interesting|great)|tell me more|send (me |over )?(your |the |some )?(more )?(info|information|details|deck|brochure|pricing)|keen to (learn|hear))\b", low):
        return "interested"
    if "?" in low and re.search(r"\b(how much|price|pricing|cost|rates?|fees?|budget|case stud(y|ies)|examples?|portfolio|references?|how (does|do|would|long)|what (is|are|would|does)|do you (have|offer|work|do)|can you|which (tech|stack|technolog))", low):
        return "question"
    return "other"


# ---------------------------------------------------------------- sending

LEGAL_SUFFIX = re.compile(r"[,\s]+(inc\.?|incorporated|llc|l\.l\.c\.|ltd\.?|limited|corp\.?|corporation|co\.|gmbh|plc|pvt\.? ltd\.?)$", re.I)


def short_company(name):
    """'Technology & Marketing Ventures, Inc.' -> 'Technology & Marketing Ventures' (reads naturally in a sentence)."""
    name = (name or "").strip()
    prev = None
    while prev != name:
        prev, name = name, LEGAL_SUFFIX.sub("", name).strip()
    return name


def render(template, lead, sender, cfg):
    ctx = SafeDict(json.loads(lead.get("extra") or "{}"))
    ctx.update(
        first_name=lead.get("first_name") or "there",
        company=short_company(lead.get("company")) or "your company",
        website=lead.get("website") or "",
        personal_line=lead.get("personal_line") or "",
        sender_name=sender["name"],
        sender_address=cfg.get("sender_address", ""),
    )
    for k, v in (cfg.get("facts") or {}).items():
        ctx.setdefault(k, v)
    text = template.format_map(ctx)
    text = "\n".join(line.rstrip() for line in text.splitlines())
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def send_email(sender, to, subject, body, in_reply_to=None, dry_run=False, cc=None, headers=None):
    try:
        import team as _team
        blocked = _team.is_trap(to) or (cc and _team.is_trap(cc))
    except Exception:
        blocked = False
    if blocked:
        raise RuntimeError("This address is on the do-not-contact list.")
    msg = EmailMessage()
    msg["From"] = f'{sender["name"]} <{sender["email"]}>'
    msg["To"] = to
    if cc:
        msg["Cc"] = cc
    for k, v in (headers or {}).items():
        msg[k] = v
    msg["Subject"] = subject
    msg["Date"] = formatdate(localtime=True)
    msg_id = make_msgid(domain=sender["email"].split("@")[1])
    msg["Message-ID"] = msg_id
    if in_reply_to:
        msg["In-Reply-To"] = in_reply_to
        msg["References"] = in_reply_to
    msg["List-Unsubscribe"] = f"<mailto:{sender['email']}?subject=unsubscribe>"
    msg.set_content(body)

    if dry_run:
        print("=" * 70)
        print(f"FROM: {msg['From']}\nTO: {to}\nSUBJECT: {subject}\n")
        print(body)
        return msg_id

    port = sender.get("smtp_port", 587)
    user = sender.get("username", sender["email"])
    if sender.get("smtp_security") == "none":  # local test mail servers only — never sends credentials
        if sender["smtp_host"] not in ("127.0.0.1", "localhost", "::1"):
            raise ValueError("Unencrypted SMTP is only allowed for a local test server")
        with smtplib.SMTP(sender["smtp_host"], port, timeout=30) as s:
            s.send_message(msg)
    elif port == 465:
        with smtplib.SMTP_SSL(sender["smtp_host"], port, timeout=30) as s:
            s.login(user, sender["password"])
            s.send_message(msg)
    else:
        with smtplib.SMTP(sender["smtp_host"], port, timeout=30) as s:
            s.starttls()
            s.login(user, sender["password"])
            s.send_message(msg)
    return msg_id


def pick_sender(lead, senders, capacity):
    if lead["sender"]:  # follow-ups stay in the same inbox/thread
        match = [s for s in senders if s["email"] == lead["sender"]]
        return match[0] if match and capacity[match[0]["email"]] > 0 else None
    open_ = [s for s in senders if capacity[s["email"]] > 0]
    return max(open_, key=lambda s: capacity[s["email"]]) if open_ else None


# ---------------------------------------------------------------- campaigns

def campaign_row(con, cid):
    return con.execute("SELECT * FROM campaigns WHERE id=?", (cid or 1,)).fetchone() or con.execute("SELECT * FROM campaigns WHERE id=1").fetchone()


def campaign_cfg(cfg, con, cid):
    """The settings one campaign runs with: its own sequence and offer if set, else the main ones."""
    c = campaign_row(con, cid)
    if not c or c["id"] == 1 and not c["sequence"]:
        return cfg
    out = dict(cfg)
    if c["sequence"]:
        try:
            seq = json.loads(c["sequence"])
            if seq:
                out["sequence"] = seq
        except ValueError:
            pass
    if c["offer"]:
        out["offer"] = c["offer"]
    return out


def campaign_inboxes(con, cid):
    c = campaign_row(con, cid)
    try:
        return {e.lower() for e in json.loads(c["inboxes"] or "[]")} if c else set()
    except ValueError:
        return set()


def pick_variant(cfg, lead, idx, force_name=None):
    """Each step can have several template variants; leads are spread evenly across the enabled ones."""
    step = cfg["sequence"][idx]
    if "variants" not in step:
        return {"name": "", "subject": step.get("subject", ""), "body": step["body"]}
    variants = [v for v in step["variants"] if v.get("enabled", True) and (v.get("body") or "").strip()]
    variants = variants or [v for v in step["variants"] if (v.get("body") or "").strip()]
    if force_name:
        for v in step["variants"]:
            if v.get("name") == force_name:
                return v
    return variants[lead["id"] % len(variants)]


def compose(cfg, con, lead, sender, idx, fresh_opener=False, variant_name=None):
    """Build (subject, body-without-footer, variant name) for sequence step idx.
    With the AI on, each email is written fresh from the campaign brief (writer.py); otherwise from the templates."""
    if not isinstance(lead, dict):
        lead = dict(lead)
    if not variant_name and not json.loads(lead.get("extra") or "{}").get("referred_by"):
        try:
            import writer
            if writer.enabled(cfg, con, lead.get("campaign_id")):
                prof = {}
                try:
                    import team
                    prof = team.state().get("strategy") or {}
                except Exception:
                    pass
                got = writer.write_step(cfg, con, lead, sender, idx, prof, len(cfg["sequence"]))
                if got:
                    return got
        except Exception as e:
            log.warning("AI writing unavailable, using the template: %s", e)
    v = pick_variant(cfg, lead, idx, variant_name)
    if idx == 0 and (fresh_opener or not lead["personal_line"]):
        lead["personal_line"] = personalize(cfg, lead)
        con.execute("UPDATE leads SET personal_line=? WHERE id=?", (lead["personal_line"], lead["id"]))
        con.commit()
    body = render(v["body"], lead, sender, cfg)
    if idx > 0 and lead["first_subject"]:
        subject = "Re: " + lead["first_subject"]  # keeps follow-ups in the same thread
    else:
        subject = render(v.get("subject") or "Quick question", lead, sender, cfg)
    x = json.loads(lead.get("extra") or "{}") if isinstance(lead, dict) else json.loads(lead["extra"] or "{}")
    if idx == 0 and x.get("referred_by"):
        rc = cfg.get("referral") or {}
        ref = SafeDict(referrer_name=x["referred_by"], referrer_first_name=x["referred_by"].split()[0])
        subject = (rc.get("subject") or "{referrer_first_name} suggested I reach out").format_map(ref)
        line = (rc.get("line") or "{referrer_name} suggested I get in touch with you.").format_map(ref)
        head, _, rest = body.partition("\n")
        body = head + "\n\n" + line + ("\n" + rest if rest else "")
    return subject, body, v.get("name", "")


def regenerate_draft(cfg, con, draft_id):
    """Rewrite a pending draft with a new AI opener (used by the CRM review page)."""
    d = con.execute("SELECT * FROM outbox WHERE id=?", (draft_id,)).fetchone()
    lead = dict(con.execute("SELECT * FROM leads WHERE id=?", (d["lead_id"],)).fetchone())
    sender = next(s for s in cfg["senders"] if s["email"] == d["sender"])
    subject, body, _ = compose(campaign_cfg(cfg, con, lead.get("campaign_id")), con, lead, sender, d["step"], fresh_opener=True, variant_name=d["variant"])
    con.execute("UPDATE outbox SET subject=?, body=? WHERE id=?", (subject, body, draft_id))
    con.commit()


def _swap_sender(body, old, new):
    """A draft moved to another inbox: sign it with the new inbox's name."""
    if old.get("name") and new.get("name") and old["name"] != new["name"]:
        return body.replace(old["name"], new["name"])
    return body


def rebalance_drafts(con, senders, capacity):
    """Waiting first-email drafts were all assigned to one inbox (older versions): spread them over the inboxes
    that have room today. Follow-ups stay with their inbox (same thread)."""
    by = {s["email"].lower(): s for s in senders}
    queued = {}
    for d in con.execute("SELECT o.id, o.sender FROM outbox o WHERE o.status IN ('pending','approved') AND o.step=0 ORDER BY o.id").fetchall():
        queued.setdefault((d["sender"] or "").lower(), []).append(d["id"])
    room = {e: capacity.get(by[e]["email"], 0) - len(queued.get(e, [])) for e in by}
    moved = 0
    for e, ids in queued.items():
        extra = len(ids) - max(0, capacity.get(by[e]["email"], 0)) if e in by else len(ids)
        for did in ids[len(ids) - extra:] if extra > 0 else []:
            d = con.execute("SELECT o.*, l.campaign_id FROM outbox o JOIN leads l ON l.id=o.lead_id WHERE o.id=?", (did,)).fetchone()
            import autopilot
            pick = {x["email"].lower() for x in autopilot.campaign_senders(con, campaign_row(con, d["campaign_id"] or 1), senders)}
            to = max((x for x in by if room.get(x, 0) > 0 and x != e and x in pick), key=lambda x: room[x], default=None)
            if not to:
                break
            old = by.get(e, {"name": ""})
            con.execute("UPDATE outbox SET sender=?, body=? WHERE id=?", (by[to]["email"], _swap_sender(d["body"], old, by[to]), did))
            room[to] -= 1
            moved += 1
    if moved:
        con.commit()
    return moved


def pick_draft_sender(lead, senders, capacity, queued):
    """For a new draft: the inbox with the most room left today once the drafts already waiting for it are counted."""
    if lead.get("sender"):
        return next((s for s in senders if s["email"] == lead["sender"]), None)
    room = {s["email"]: capacity.get(s["email"], 0) - queued.get(s["email"].lower(), 0) for s in senders}
    open_ = [s for s in senders if room[s["email"]] > 0]
    return max(open_, key=lambda s: room[s["email"]]) if open_ else None


def _autopilot_round(cfg, con, args, due, for_campaign, usable, capacity, limits, allowance, all_senders, master_approves, redirect,
                     queued, draft_budget, cap_left):
    """One round of the BSO Bot, built for hundreds of inboxes: decide what each inbox does (sequential, fast), write the
    emails in parallel, send them in parallel (each inbox sends at most one email per round), then record the results.
    Returns (sent, drafted, waiting, failed, queued_first, held, problems)."""
    import autopilot
    import convo
    from concurrent.futures import ThreadPoolExecutor
    ap = autopilot.conf(cfg)
    timed = not args.force
    sent = drafted = waiting = failed = queued_first = held = 0
    problems, plan, stocks = [], [], {}
    planned_first, planned_camp, planned_dom = {}, {}, {}

    def dom_ok(s):
        d = autopilot.domain_of(s["email"])
        r = con.execute("SELECT cap FROM sending_domains WHERE domain=?", (d,)).fetchone()
        return not (r and r["cap"] and autopilot.sent_today_domain(con, d) + planned_dom.get(d, 0) >= r["cap"])

    def ok_now(s):
        return usable(s) and dom_ok(s)

    def first_room(cid, c):
        if cid not in stocks:
            stocks[cid] = autopilot.new_today_quota(autopilot.stock(cfg, con, c, limits, all_senders, allowance))
        return (autopilot.first_emails_today(con, cid) + planned_first.get(cid, 0) < stocks[cid] and
                autopilot.sent_today_campaign(con, cid) + planned_camp.get(cid, 0) < allowance.get(cid, 0))

    def reserve(s, cid, idx):
        capacity[s["email"]] -= 1
        autopilot.after_send(cfg, con, s["email"])              # the inbox's 6-20 minute rest starts now
        d = autopilot.domain_of(s["email"])
        planned_dom[d] = planned_dom.get(d, 0) + 1
        planned_camp[cid] = planned_camp.get(cid, 0) + 1
        if idx == 0:
            planned_first[cid] = planned_first.get(cid, 0) + 1

    for row in due:
        lead = dict(row)
        idx = lead["step"]
        if convo.protected(con, lead["email"]):
            print(f"SKIP: {lead['email']}: part of an open conversation with their company; emails paused")
            con.execute("UPDATE leads SET status='paused' WHERE id=?", (lead["id"],))
            con.execute("INSERT INTO notes(lead_id, body, at, kind) VALUES (?,?,?,'note')",
                        (lead["id"], "Emails paused: this person is part of an open conversation with their company.", iso(now_utc())))
            continue
        cid = lead.get("campaign_id") or 1
        ccfg, senders, c = for_campaign(cid)
        steps = ccfg["sequence"]
        if not senders:
            problems.append(f"Campaign {c['name'] if c else cid}: none of its inboxes can send right now.")
            continue
        if idx >= len(steps):
            con.execute("UPDATE leads SET status='done' WHERE id=?", (lead["id"],))
            continue
        if timed and idx == 0:
            why, until = autopilot.first_email_hold(cfg, con, lead, c)
            if why:
                nxt = autopilot.plan_arrival(cfg, c, lead, until)
                con.execute("UPDATE leads SET next_send_at=? WHERE id=?", (iso(nxt), lead["id"]))
                autopilot.note(con, "hold", why.split(" on ")[0] if " on " in why else re.sub(r"\d+", "N", why), lead["id"], cid)
                held += 1
                continue
        draft = con.execute("SELECT * FROM outbox WHERE lead_id=? AND step=? AND status IN ('pending','approved')", (lead["id"], idx)).fetchone()
        if draft and draft["status"] == "pending":
            waiting += 1
            continue
        will_draft = not draft and (master_approves or autopilot.autonomy_of(c) == "assisted")
        if timed and not will_draft:                            # it goes out now: only when it may land
            if idx == 0 and not autopilot.in_window(cfg, c, lead):
                con.execute("UPDATE leads SET next_send_at=? WHERE id=?", (iso(autopilot.plan_arrival(cfg, c, lead)), lead["id"]))
                queued_first += 1
                continue
            if idx > 0 and not autopilot.in_hours(cfg, lead):
                continue
            if idx == 0 and not first_room(cid, c):
                continue
        if cap_left is not None and sum(planned_camp.values()) >= cap_left:
            break
        if draft:                                                # approved: send it from an inbox that may send now
            sender = next((s for s in senders if s["email"] == draft["sender"]), None)
            if idx == 0 and (not sender or not ok_now(sender)):
                other = next((s for s in sorted(senders, key=lambda s: -capacity.get(s["email"], 0)) if ok_now(s)), None)
                if other:
                    con.execute("UPDATE outbox SET sender=?, body=? WHERE id=?", (other["email"], _swap_sender(draft["body"], sender or {}, other), draft["id"]))
                    draft = con.execute("SELECT * FROM outbox WHERE id=?", (draft["id"],)).fetchone()
                    sender = other
            if not sender or not ok_now(sender):
                continue
            reserve(sender, cid, idx)
            plan.append({"lead": lead, "c": c, "ccfg": ccfg, "steps": steps, "idx": idx, "sender": sender, "draft": draft, "alone": False,
                         "subject": draft["subject"], "body": draft["body"], "variant": draft["variant"], "score": draft["quality"]})
            continue
        if will_draft:
            if draft_budget <= 0:
                continue
            pool = [s for s in senders if capacity.get(s["email"], 0) > 0]
            sender = pick_draft_sender(lead, pool if not lead.get("sender") else senders, capacity, queued)
            if not sender:
                continue
            queued[sender["email"].lower()] = queued.get(sender["email"].lower(), 0) + 1
            draft_budget -= 1
        else:
            pool = [s for s in senders if ok_now(s)]
            sender = pick_draft_sender(lead, pool if not lead.get("sender") else senders, capacity, queued)
            if not sender or not ok_now(sender):
                continue
            reserve(sender, cid, idx)
        plan.append({"lead": lead, "c": c, "ccfg": ccfg, "steps": steps, "idx": idx, "sender": sender, "draft": None,
                     "will_draft": will_draft, "compose": True})
    con.commit()

    def write(p):
        tcon = db_connect(cfg.get("database", "outreach.db"))
        try:
            subject, body, variant = compose(p["ccfg"], tcon, p["lead"], p["sender"], p["idx"])
            hits = find_avoid(p["ccfg"], subject, body)
            score, why = autopilot.quality(cfg, tcon, p["ccfg"], p["lead"], subject, body, ai="|" in (variant or ""), sender_name=p["sender"].get("name", ""))
            return dict(subject=subject, body=body, variant=variant, hits=hits, score=score, why=why)
        except Exception as e:
            return dict(error=str(e))
        finally:
            tcon.close()
    jobs = [p for p in plan if p.get("compose")]
    if jobs:
        with ThreadPoolExecutor(max_workers=max(1, int(ap.get("write_workers") or 6))) as ex:
            for p, out in zip(jobs, ex.map(write, jobs)):
                p.update(out)
    to_send = []
    for p in plan:
        lead, idx, sender, c = p["lead"], p["idx"], p["sender"], p["c"]
        if p.get("error"):
            problems.append(f"Couldn't write the email for {lead['email']}: {p['error']}")
            continue
        if p.get("hits"):
            problems.append(f"Step {idx + 1} email for {lead['email']} uses avoided phrases ({', '.join(p['hits'])}); not drafted or sent.")
            continue
        if p.get("compose"):
            if p["will_draft"] or not autopilot.may_send_alone(cfg, c, p["score"]):
                con.execute("INSERT INTO outbox(lead_id, step, sender, subject, body, status, created_at, variant, quality, quality_notes, orig_body) "
                            "VALUES (?,?,?,?,?,'pending',?,?,?,?,?)",
                            (lead["id"], idx, sender["email"], p["subject"], p["body"], iso(now_utc()), p["variant"], p["score"],
                             "; ".join(p["why"]) or None, p["body"]))
                log_event(con, lead["id"], "drafted", f"step {idx + 1} (quality {p['score']})", sender["email"])
                drafted += 1
                continue
            p["alone"] = True
        full_body = p["body"] + "\n\n" + render(p["ccfg"]["footer"], lead, sender, p["ccfg"])
        rc = cfg.get("ref_code") or {}
        if rc.get("enabled", True):
            code = convo.ref_code(con, lead["id"], rc.get("prefix") or "NS")
            if code:
                full_body += f"\nRef: {code}"
        p["full_body"] = full_body
        p["to"], p["shown"] = (redirect, f"[TEST for {lead['email']}] {p['subject']}") if redirect else (lead["email"], p["subject"])
        to_send.append(p)
    con.commit()

    def deliver(p):
        try:
            return send_email(p["sender"], p["to"], p["shown"], p["full_body"], p["lead"]["first_msg_id"] if p["idx"] > 0 else None, False), None
        except Exception as e:
            return None, e
    if to_send:
        with ThreadPoolExecutor(max_workers=max(1, int(ap.get("send_workers") or 12))) as ex:
            results = list(ex.map(deliver, to_send))
    else:
        results = []
    for p, (msg_id, err) in zip(to_send, results):
        lead, idx, sender, steps = p["lead"], p["idx"], p["sender"], p["steps"]
        if err is not None:
            if isinstance(err, smtplib.SMTPAuthenticationError):
                guard_login(cfg, con, sender["email"], False)
                problems.append(f"{sender['email']}: login rejected. Check its password.")
            elif isinstance(err, (ConnectionError, TimeoutError, socket.gaierror)):
                problems.append(f"{sender['email']}: can't reach the mail server ({err}).")
            else:
                log_event(con, lead["id"], "error", str(err), sender["email"])
                failed += 1
            continue
        convo.record_sent(con, msg_id, lead, "step", sender["email"], p["to"], p["shown"], p["full_body"], step=idx, variant=p["variant"],
                          test=bool(redirect), angle=p["variant"] if "|" in (p["variant"] or "") else None)
        sent += 1
        guard_login(cfg, con, sender["email"], True)
        if p["alone"] and autopilot.maybe_spot_check(cfg, con, p["c"], lead, sender["email"], p["subject"], p["body"], p["score"] or 0):
            print(f"Spot check: the email to {lead['email']} was set aside for a person to look at.")
        if p["draft"]:
            con.execute("UPDATE outbox SET status='sent', sent_at=? WHERE id=?", (iso(now_utc()), p["draft"]["id"]))
        nxt = idx + 1
        next_at, status = (iso(now_utc() + timedelta(days=steps[nxt]["delay_days"])), "active") if nxt < len(steps) else (None, "done")
        con.execute("UPDATE leads SET step=?, sender=?, next_send_at=?, status=?, first_msg_id=COALESCE(first_msg_id, ?), "
                    "first_subject=COALESCE(first_subject, ?) WHERE id=?", (nxt, sender["email"], next_at, status, msg_id, p["subject"], lead["id"]))
        log_event(con, lead["id"], "sent", f"step {idx + 1}" + (" (test)" if redirect else "") + (" (sent on its own)" if p["alone"] else ""),
                  sender["email"], p["variant"])
        if lead["stage"] == "new":
            set_stage(con, lead["id"], "contacted")
    con.commit()
    return sent, drafted, waiting, failed, queued_first, held, problems


def cmd_run(cfg, con, args):
    import autopilot
    import convo
    lg_ver = str((cfg.get("legal") or {}).get("version", "1.0"))
    if str((cfg.get("legal_acceptance") or {}).get("version")) != lg_ver and not args.dry_run:
        print("PROBLEM: Open BSO and accept the Terms of Use first.")
        sys.exit(1)
    team_cap = 0
    team_mod = None
    if not args.dry_run:
        import team as team_mod
        if team_mod.joined() and not team_mod.is_bot() and (team_mod.policy() or {}).get("role"):
            print("Cold email is sent only by the BSO Bot. This BSO handles your own prospects and conversations; "
                  "the cold inboxes belong on the bot (Master Console → BSO setup).")
            return
        ok, why, team_cap = team_mod.gate(con)
        if not ok:
            print("PROBLEM:", why)
            sys.exit(1)
    ap_on = autopilot.active(cfg) and not args.dry_run
    if ap_on:
        try:
            import ops
            for line in ops.housekeeping(cfg, con):
                print(line)
        except Exception as e:
            print(f"PROBLEM: housekeeping failed ({e}).")
        k, kwhy = autopilot.killed(con)
        if k:
            print("Sending is stopped. " + kwhy)
            return
    elif not args.force and not in_send_window(cfg):
        print("Outside the sending window; nothing sent (use --force to override).")
        return
    if not args.dry_run and not args.no_check:
        last = con.execute("SELECT v FROM app_mode WHERE k='last_check'").fetchone()
        if not ap_on or args.force or not last or last[0] < iso(now_utc() - timedelta(minutes=15)):
            cmd_check(cfg, con, args)          # replies arrive live; the full check is a safety net every 15 minutes
            con.execute("INSERT INTO app_mode(k, v) VALUES ('last_check', ?) ON CONFLICT(k) DO UPDATE SET v=excluded.v", (iso(now_utc()),))
            con.commit()

    growth = []
    if not args.dry_run:
        for modname in ("instantly", "discover"):
            try:
                mod = __import__(modname)
                note = mod.maybe_pull(cfg, con) if modname == "instantly" else mod.maybe_discover(cfg, con)
                if note:
                    growth.append(note)
            except Exception as e:
                growth.append(f"PROBLEM: {modname} step failed ({e}).")
        for g in growth:
            print(g)

    senders = [s for s in load_senders(cfg) if s["password"] or args.dry_run]
    guard_notes = [] if args.dry_run else guard_check(cfg, con, senders)
    paused_inboxes = guard_paused(con)
    if paused_inboxes and not args.dry_run:
        senders = [s for s in senders if s["email"].lower() not in paused_inboxes]
        for e, why in paused_inboxes.items():
            print(f"Inbox {e} is paused by the deliverability guard: {why}")
    if ap_on:                                       # today's limit follows each inbox's reputation
        for s in senders:
            s["daily_limit"], s["limit_why"] = autopilot.limit_for(cfg, con, s)
        senders = [s for s in senders if s["daily_limit"] > 0] or senders[:0]
    if not senders:
        print("No usable sender inboxes (check your inbox passwords, and that each inbox passed its readiness check).")
        return
    capacity = {s["email"]: s["daily_limit"] - sent_today(con, s["email"], cfg) for s in senders}
    limits = {s["email"].lower(): s["daily_limit"] for s in senders}

    review = ((cfg.get("review") or {}).get("enabled", False) or (team_mod and team_mod.holds_approval())) and not args.dry_run
    master_approves = bool(team_mod and team_mod.holds_approval())
    cap_left = None
    if team_cap:
        cap_left = max(0, team_cap - team_mod.sent_today_total(con))
        if cap_left == 0:
            print(f"Your manager's daily cap ({team_cap}) is reached; nothing more is sent today.")
            return
    tm = cfg.get("test_mode") or {}
    redirect = tm.get("redirect_to") if tm.get("enabled") else None
    if tm.get("enabled") and not redirect:
        print("test_mode is on but redirect_to is empty; refusing to send.")
        return

    if not args.dry_run:
        moved = rebalance_drafts(con, senders, capacity)
        if moved:
            print(f"Moved {moved} waiting draft(s) to inboxes with room today, so every inbox sends its share.")
    queued = {}
    for r in con.execute("SELECT lower(sender) e, COUNT(*) n FROM outbox WHERE status IN ('pending','approved') GROUP BY lower(sender)"):
        queued[r["e"]] = r["n"]
    draft_budget = max(0, sum(max(c, 0) for c in capacity.values()) - sum(queued.values()))
    if cap_left is not None:
        draft_budget = min(draft_budget, max(0, cap_left - sum(queued.values())))

    try:                                     # lead scoring from the approved strategy: the most promising new leads go first
        import learn
        import writer
        weights = ((team_mod.state() if team_mod else {}).get("strategy") or {}).get("scoring")
        if weights:
            for r in con.execute("SELECT id, extra FROM leads WHERE status='active' AND step=0 AND fit_score IS NULL LIMIT 5000").fetchall():
                con.execute("UPDATE leads SET fit_score=? WHERE id=?", (learn.lead_score(weights, r["extra"], writer.segment(dict(r))), r["id"]))
            con.commit()
    except Exception as e:
        log.warning("Lead scoring skipped: %s", e)
    due = con.execute(
        "SELECT l.* FROM leads l LEFT JOIN campaigns c ON c.id=COALESCE(l.campaign_id,1) "
        "WHERE l.status='active' AND l.next_send_at<=? AND COALESCE(c.status,'active')='active' "
        "ORDER BY l.step DESC, COALESCE(l.fit_score, 50) DESC, l.id LIMIT ?",  # running sequences first, then the likeliest new leads; paused campaigns wait
        (iso(now_utc()), args.limit if not ap_on else max(args.limit, 400, 3 * len(senders))),
    ).fetchall()
    all_senders, ccache, stocks = senders, {}, {}
    allowance = autopilot.campaign_allowance(cfg, con, limits, all_senders) if ap_on else {}

    def for_campaign(cid):
        if cid not in ccache:
            c = campaign_row(con, cid)
            if ap_on:
                pool = autopilot.campaign_senders(con, c, all_senders)
            else:
                pick = campaign_inboxes(con, cid)
                pool = [s for s in all_senders if not pick or s["email"].lower() in pick]
            ccache[cid] = (campaign_cfg(cfg, con, cid), pool, c)
        return ccache[cid]

    def usable(s):
        if capacity.get(s["email"], 0) <= 0:
            return False
        if ap_on and not args.force:
            if not autopilot.gap_ok(con, s["email"]):
                return False
            d = con.execute("SELECT cap FROM sending_domains WHERE domain=?", (autopilot.domain_of(s["email"]),)).fetchone()
            if d and d["cap"] and autopilot.sent_today_domain(con, autopilot.domain_of(s["email"])) >= d["cap"]:
                return False
        return True

    sent = drafted = waiting = failed = queued_first = held = 0
    problems = []
    lo, hi = cfg.get("delay_between_sends_seconds", [60, 180])
    if ap_on:
        lo, hi = 3, 12                       # different inboxes in turn; each inbox keeps its own 6-20 minute gap
    if ap_on:
        sent, drafted, waiting, failed, queued_first, held, problems = _autopilot_round(
            cfg, con, args, due, for_campaign, usable, capacity, limits, allowance, all_senders, master_approves, redirect, queued, draft_budget,
            cap_left)
    else:
        for row in due:
            lead = dict(row)
            idx = lead["step"]
            if convo.protected(con, lead["email"]):                # part of a live deal elsewhere: no cold emails
                print(f"SKIP: {lead['email']}: part of an open conversation with their company; emails paused")
                con.execute("UPDATE leads SET status='paused' WHERE id=?", (lead["id"],))
                con.execute("INSERT INTO notes(lead_id, body, at, kind) VALUES (?,?,?,'note')",
                            (lead["id"], "Emails paused: this person is part of an open conversation with their company.", iso(now_utc())))
                continue
            cid = lead.get("campaign_id") or 1
            ccfg, senders, c = for_campaign(cid)
            steps = ccfg["sequence"]
            if not senders:
                problems.append(f"Campaign {c['name'] if c else cid}: none of its inboxes can send right now.")
                continue
            if idx >= len(steps):
                con.execute("UPDATE leads SET status='done' WHERE id=?", (lead["id"],))
                continue

            timed = ap_on and not args.force
            if timed and idx == 0:
                why, until = autopilot.first_email_hold(cfg, con, lead, c)
                if why:
                    nxt = autopilot.plan_arrival(cfg, c, lead, until)
                    con.execute("UPDATE leads SET next_send_at=? WHERE id=?", (iso(nxt), lead["id"]))
                    autopilot.note(con, "hold", why.split(" on ")[0] if " on " in why else re.sub(r"\d+", "N", why), lead["id"], cid)
                    print(f"HOLD: {lead['email']}: {why}; next try {iso(nxt)[:16]}")
                    held += 1
                    continue

            draft = con.execute("SELECT * FROM outbox WHERE lead_id=? AND step=? AND status IN ('pending','approved')", (lead["id"], idx)).fetchone() \
                if not args.dry_run else None
            alone = False
            if draft and draft["status"] == "pending":
                waiting += 1
                continue
            if draft:                                             # approved: send it, from an inbox with room
                sender = next((s for s in senders if s["email"] == draft["sender"]), None)
                if idx == 0 and (not sender or not usable(sender)):
                    other = next((s for s in sorted(senders, key=lambda s: -capacity.get(s["email"], 0)) if usable(s)), None)
                    if other:
                        con.execute("UPDATE outbox SET sender=?, body=? WHERE id=?", (other["email"], _swap_sender(draft["body"], sender or {}, other), draft["id"]))
                        draft = con.execute("SELECT * FROM outbox WHERE id=?", (draft["id"],)).fetchone()
                        sender = other
                if not sender or not usable(sender):
                    continue
                subject, body, variant = draft["subject"], draft["body"], draft["variant"]
                score = draft["quality"] if "quality" in draft.keys() else None
            else:
                will_draft = (master_approves or autopilot.autonomy_of(c) == "assisted") if ap_on else review
                if timed and not will_draft:                      # it would go straight out: only compose it when it may land
                    if idx == 0 and not autopilot.in_window(cfg, c, lead):
                        con.execute("UPDATE leads SET next_send_at=? WHERE id=?", (iso(autopilot.plan_arrival(cfg, c, lead)), lead["id"]))
                        queued_first += 1
                        continue
                    if idx > 0 and not autopilot.in_hours(cfg, lead):
                        continue
                    if idx == 0:
                        if cid not in stocks:
                            stocks[cid] = autopilot.new_today_quota(autopilot.stock(cfg, con, c, limits, all_senders, allowance))
                        if autopilot.first_emails_today(con, cid) >= stocks[cid] or autopilot.sent_today_campaign(con, cid) >= allowance.get(cid, 0):
                            continue
                if review or ap_on:
                    if will_draft:
                        pool = [s for s in senders if capacity.get(s["email"], 0) > 0]
                    else:
                        pool = [s for s in senders if usable(s)]
                    sender = pick_draft_sender(lead, pool if not lead.get("sender") else senders, capacity, queued)
                    if lead.get("sender") and sender and ap_on and not will_draft and not usable(sender):
                        continue                                   # the follow-up's inbox is resting or full: next tick
                else:
                    sender = pick_sender(lead, senders, capacity)
                if not sender:
                    continue
                subject, body, variant = compose(ccfg, con, lead, sender, idx)
                hits = find_avoid(ccfg, subject, body)
                if hits:
                    problems.append(f"Step {idx + 1} email for {lead['email']} uses avoided phrases ({', '.join(hits)}); not drafted or sent. "
                                    "Edit the template in Settings → Your offer & emails.")
                    continue
                score, why_q = autopilot.quality(cfg, con, ccfg, lead, subject, body, ai="|" in (variant or ""), sender_name=sender.get("name", "")) if ap_on else (None, [])
                alone = ap_on and not master_approves and autopilot.may_send_alone(cfg, c, score) if ap_on else not review
                if not alone:
                    if draft_budget <= 0:
                        continue
                    con.execute(
                        "INSERT INTO outbox(lead_id, step, sender, subject, body, status, created_at, variant, quality, quality_notes, orig_body) "
                        "VALUES (?,?,?,?,?,'pending',?,?,?,?,?)",
                        (lead["id"], idx, sender["email"], subject, body, iso(now_utc()), variant, score, "; ".join(why_q) or None, body))
                    log_event(con, lead["id"], "drafted", f"step {idx + 1}" + (f" (quality {score})" if score is not None else ""), sender["email"])
                    queued[sender["email"].lower()] = queued.get(sender["email"].lower(), 0) + 1
                    con.commit()
                    drafted += 1
                    draft_budget -= 1
                    continue

            if timed:                                             # when it may land
                if idx == 0:
                    if not autopilot.in_window(cfg, c, lead):
                        nxt = autopilot.plan_arrival(cfg, c, lead)
                        con.execute("UPDATE leads SET next_send_at=? WHERE id=?", (iso(nxt), lead["id"]))
                        queued_first += 1
                        continue
                    if cid not in stocks:
                        stocks[cid] = autopilot.new_today_quota(autopilot.stock(cfg, con, c, limits, all_senders, allowance))
                    if autopilot.first_emails_today(con, cid) >= stocks[cid] or autopilot.sent_today_campaign(con, cid) >= allowance.get(cid, 0):
                        continue                                   # today's new starts are done; capacity is kept for follow-ups
                elif not autopilot.in_hours(cfg, lead):
                    continue

            reply_to = lead["first_msg_id"] if idx > 0 else None
            full_body = body + "\n\n" + render(ccfg["footer"], lead, sender, ccfg)
            rc = cfg.get("ref_code") or {}
            if rc.get("enabled", True):
                code = convo.ref_code(con, lead["id"], rc.get("prefix") or "NS")
                if code:
                    full_body += f"\nRef: {code}"
            to, shown_subject = lead["email"], subject
            if redirect:  # test mode: real lead data, but it lands in your own inbox
                to, shown_subject = redirect, f"[TEST for {lead['email']}] {subject}"

            if cap_left is not None and sent >= cap_left:
                break
            if sent > 0 and not args.dry_run:
                time.sleep(random.uniform(lo, hi))  # human-ish pacing
            try:
                msg_id = send_email(sender, to, shown_subject, full_body, reply_to, args.dry_run)
            except smtplib.SMTPAuthenticationError as e:
                log.error("Login failed for %s: %s — skipping this inbox", sender["email"], e)
                capacity[sender["email"]] = 0
                if not args.dry_run:
                    guard_login(cfg, con, sender["email"], False)
                problems.append(f"{sender['email']}: login rejected. Check the app password in Settings → Email inboxes.")
                continue
            except (ConnectionError, TimeoutError, socket.gaierror) as e:
                log.error("Can't reach the mail server for %s: %s — skipping this inbox", sender["email"], e)
                capacity[sender["email"]] = 0
                problems.append(f"{sender['email']}: can't reach the mail server ({e}). Check the server details or your internet connection.")
                continue
            except Exception as e:
                log.error("Send to %s failed: %s", lead["email"], e)
                log_event(con, lead["id"], "error", str(e), sender["email"])
                con.commit()
                failed += 1
                continue

            convo.record_sent(con, msg_id, lead, "step", sender["email"], to, shown_subject, full_body, step=idx, variant=variant,
                              test=bool(redirect), angle=variant if "|" in (variant or "") else None)
            sent += 1
            capacity[sender["email"]] -= 1
            if not args.dry_run:
                guard_login(cfg, con, sender["email"], True)
            if args.dry_run:
                continue
            if ap_on:
                autopilot.after_send(cfg, con, sender["email"])
                if alone and autopilot.maybe_spot_check(cfg, con, c, lead, sender["email"], subject, body, score or 0):
                    print(f"Spot check: the email to {lead['email']} was set aside for a person to look at.")
            if draft:
                con.execute("UPDATE outbox SET status='sent', sent_at=? WHERE id=?", (iso(now_utc()), draft["id"]))

            nxt = idx + 1
            if nxt < len(steps):
                next_at = iso(now_utc() + timedelta(days=steps[nxt]["delay_days"]))
                status = "active"
            else:
                next_at, status = None, "done"
            con.execute(
                "UPDATE leads SET step=?, sender=?, next_send_at=?, status=?, "
                "first_msg_id=COALESCE(first_msg_id, ?), first_subject=COALESCE(first_subject, ?) "
                "WHERE id=?",
                (nxt, sender["email"], next_at, status, msg_id, subject, lead["id"]),
            )
            log_event(con, lead["id"], "sent", f"step {idx + 1}" + (" (test)" if redirect else "") + (" (sent on its own)" if alone and ap_on else ""),
                      sender["email"], variant)
            if lead["stage"] == "new":
                set_stage(con, lead["id"], "contacted")
            con.commit()
            log.info("Sent step %d to %s via %s", idx + 1, to, sender["email"])

    con.commit()
    if not args.dry_run:
        try:
            r_drafted, r_sent = nurture_tick(cfg, con, all_senders, redirect)
            if r_drafted:
                print(f"Drafted {r_drafted} check-back email(s) for leads who asked us to come back later: waiting for review.")
            sent += r_sent
        except Exception as e:
            problems.append(f"Check-back emails: {e}")
    if team_mod:
        try:
            team_mod.sync(con)
        except Exception as e:  # reporting must never break sending
            log.warning("Team sync failed: %s", e)
    msg = f"{'Previewed' if args.dry_run else 'Sent'} {sent} email(s)."
    if review or drafted or waiting:
        msg += f" Drafted {drafted} new for review; {waiting} waiting for your approval in the CRM."
    if queued_first:
        msg += f" {queued_first} first email(s) queued for their arrival window."
    if held:
        msg += f" {held} held back by the contact rules."
    if failed:
        msg += f" {failed} failed to send (details above)."
    print(msg)
    for p in dict.fromkeys(guard_notes + problems):
        print("PROBLEM:", p)
    if problems or failed:
        sys.exit(1)  # marks the run as Failed in the CRM so it's noticed


DEFAULT_RECONNECT = """Hi {first_name},

When we last spoke you mentioned it could make sense to pick this up again around now, so I'm checking in.

Is {company} still thinking about extra engineering capacity? Happy to share what's changed on our side, or to leave it if the timing still isn't right.

Best,
{sender_name}"""


def reconnect_draft(cfg, con, lead, why="check back"):
    """A soft check-back / reconnect email, always waiting for a human to review it."""
    lead = dict(lead)
    have = con.execute("SELECT 1 FROM outbox WHERE lead_id=? AND kind='reconnect' AND status IN ('pending','approved')", (lead["id"],)).fetchone()
    if have:
        return None
    senders = cfg.get("senders") or []
    sender = next((x for x in senders if x.get("email", "").lower() == (lead.get("sender") or "").lower()), None) or (senders[0] if senders else None)
    if not sender:
        return None
    body = None
    try:
        import writer
        body = writer.reconnect(cfg, con, lead, sender, why)
    except Exception as e:
        log.warning("AI check-back draft unavailable: %s", e)
    body = body or render((cfg.get("nurture") or {}).get("template") or DEFAULT_RECONNECT, lead, sender, cfg)
    subject = "Re: " + (lead.get("first_subject") or "following up") if lead.get("first_subject") else f"Following up, {lead.get('first_name') or ''}".strip(", ")
    cur = con.execute("INSERT INTO outbox(lead_id, step, sender, subject, body, status, created_at, kind) VALUES (?,?,?,?,?,'pending',?,'reconnect')",
                      (lead["id"], -1, sender["email"], subject, body, iso(now_utc())))
    log_event(con, lead["id"], "drafted", f"check-back email ({why})", sender["email"])
    return cur.lastrowid


def nurture_tick(cfg, con, senders, redirect):
    """Leads who said 'not now': draft the check-back email when their date comes; send the ones a person approved."""
    import convo
    today = date.today().isoformat()
    drafted = sent = 0
    for l in con.execute("SELECT * FROM leads WHERE status='nurture' AND check_back IS NOT NULL AND check_back<=?", (today,)).fetchall():
        if reconnect_draft(cfg, con, l, "they asked us to check back"):
            drafted += 1
    by_email = {x["email"].lower(): x for x in senders}
    for d in con.execute("SELECT o.*, l.email lemail, l.first_msg_id FROM outbox o JOIN leads l ON l.id=o.lead_id "
                         "WHERE o.kind='reconnect' AND o.status='approved'").fetchall():
        sender = by_email.get((d["sender"] or "").lower()) or (senders[0] if senders else None)
        if not sender:
            continue
        lead = dict(con.execute("SELECT * FROM leads WHERE id=?", (d["lead_id"],)).fetchone())
        if lead["status"] in ("unsubscribed", "bounced") or con.execute("SELECT 1 FROM suppression WHERE email=?", (lead["email"],)).fetchone():
            con.execute("UPDATE outbox SET status='cancelled' WHERE id=?", (d["id"],))
            continue
        to, subj = (redirect, f"[TEST for {lead['email']}] {d['subject']}") if redirect else (lead["email"], d["subject"])
        body = d["body"] + "\n\n" + render(cfg.get("footer") or "", lead, sender, cfg)
        mid = send_email(sender, to, subj, body, lead.get("first_msg_id"))
        convo.record_sent(con, mid, lead, "reconnect", sender["email"], to, subj, body, test=bool(redirect))
        con.execute("UPDATE outbox SET status='sent', sent_at=? WHERE id=?", (iso(now_utc()), d["id"]))
        con.execute("UPDATE leads SET check_back=NULL WHERE id=?", (lead["id"],))
        log_event(con, lead["id"], "sent", "check-back email" + (" (test)" if redirect else ""), sender["email"])
        sent += 1
    con.commit()
    return drafted, sent


# ---------------------------------------------------------------- inbox

def get_text_body(msg):
    parts = msg.walk() if msg.is_multipart() else [msg]
    plain = html = ""
    for p in parts:
        if p.get_content_disposition() == "attachment":
            continue
        ctype = p.get_content_type()
        if ctype not in ("text/plain", "text/html"):
            continue
        try:
            text = p.get_payload(decode=True).decode(p.get_content_charset() or "utf-8", "replace")
        except Exception:
            continue
        if ctype == "text/plain" and not plain:
            plain = text
        elif ctype == "text/html" and not html:
            html = re.sub(r"<[^>]+>", " ", text)
    text = plain or html
    # keep only the new part of the reply, drop the quoted thread
    lines = []
    for line in text.splitlines():
        if line.startswith(">") or re.match(r"^On .+wrote:\s*$", line.strip()):
            break
        lines.append(line)
    return "\n".join(lines).strip()


def notify(text, cfg=None):
    print(text)
    to = ((cfg or {}).get("notifications") or {}).get("email", "").strip()
    if to and cfg:
        try:  # an email alert, sent from the first working inbox
            sender = next((x for x in load_senders(cfg) if x["password"]), None)
            if sender:
                send_email(sender, to, "BSO: " + text.split(":", 1)[0][:80], text)
        except Exception as e:
            log.warning("Email notification failed: %s", e)
    hook = env("SLACK_WEBHOOK_URL")
    if hook:
        try:
            requests.post(hook, json={"text": text}, timeout=10)
        except Exception as e:
            log.warning("Slack notify failed: %s", e)


# ---------------------------------------------------------------- leads by email

INTAKE_TYPES = (".csv", ".xlsx", ".xlsm")
INTAKE_MAX_BYTES = 15 * 1024 * 1024


def _sender_verified(msg, addr):
    """True when the receiving mail server confirmed the sender (DMARC, or DKIM signed by the sender's own domain)."""
    ar = " ".join(str(h) for h in (msg.get_all("Authentication-Results") or [])).lower()
    dom = addr.split("@")[-1].lower()
    return "dmarc=pass" in ar or ("dkim=pass" in ar and f"header.d={dom}" in ar)


def intake_message(cfg, con, sender, msg):
    """A lead file emailed to the bot's intake inbox by an allowed person: import it and reply with the result.
    Returns True when the message was an intake email (handled here), False otherwise."""
    li = cfg.get("lead_intake") or {}
    allowed = {a.strip().lower() for a in li.get("allowed_senders") or [] if a.strip()}
    if not li.get("enabled") or not allowed or sender["email"].lower() != str(li.get("inbox", "")).lower():
        return False
    frm = parseaddr(msg.get("From", ""))[1].lower()
    if frm not in allowed:
        return False
    files = []
    for part in msg.walk():
        name = part.get_filename() or ""
        if name.lower().endswith(INTAKE_TYPES):
            files.append((name, part.get_payload(decode=True) or b""))
    subject = msg.get("Subject", "") or "Leads"
    def answer(text):
        try:
            send_email(sender, frm, subject if subject.lower().startswith("re:") else "Re: " + subject, text, msg.get("Message-ID"),
                       headers={"X-BSO-Kind": "intake"})
        except Exception as e:
            log.error("Couldn't answer the lead file email: %s", e)
    if not files and not re.search(r"\b(leads?|campaign)\b", subject, re.I):
        return False          # an ordinary email (a reply, a question): not a lead file, let normal handling see it
    if not files:
        answer("BSO found no lead file in your email. Attach a CSV or Excel (.xlsx) file with one lead per row and an email column, and send it again.")
        return True
    if li.get("verified_only", True) and not _sender_verified(msg, frm):
        log.warning("Lead file from %s refused: sender not verified by the mail server", frm)
        answer("BSO did not import this file: the mail server couldn't confirm it really came from your address (no DMARC/DKIM pass). "
               "If you sent it, ask your admin to switch off \"verified senders only\" in BSO → Settings → Lead targeting, or send from your company mailbox.")
        notify(f"⚠ A lead file emailed from {frm} was refused because the sender couldn't be verified.", cfg)
        return True
    status = "active" if li.get("add_to_sequence", True) else "paused"
    cid, cname = int(li.get("campaign_id") or 1), None
    m = re.search(r"campaign\s*[:=-]\s*(.+)$", subject, re.I)        # "Leads, campaign: UK SaaS" picks a campaign by name
    if m:
        want = m.group(1).strip().lower()
        hit = con.execute("SELECT id, name FROM campaigns WHERE lower(name)=?", (want,)).fetchone() or \
            con.execute("SELECT id, name FROM campaigns WHERE lower(name) LIKE ?", (want + "%",)).fetchone()
        if hit:
            cid, cname = hit["id"], hit["name"]
    cname = cname or campaign_row(con, cid)["name"]
    lines, total = [], 0
    for name, data in files:
        if len(data) > INTAKE_MAX_BYTES:
            lines.append(f"{name}: too large (over 15 MB), not imported.")
            continue
        try:
            if name.lower().endswith(".csv"):
                try:
                    text = data.decode("utf-8-sig")
                except UnicodeDecodeError:
                    text = data.decode("latin-1")
                rows = read_rows(io.StringIO(text), name)
            else:
                rows = read_rows(io.BytesIO(data), name)
            added, skipped = import_rows(con, rows, status, source=f"emailed by {frm}: {name}", campaign_id=cid)
        except Exception as e:
            lines.append(f"{name}: couldn't be read ({e}). Check it's a normal CSV or Excel file with a header row.")
            continue
        total += added
        why = ", ".join(f"{n} {r}" for r, n in sorted(skipped.items(), key=lambda x: -x[1])) or "none"
        lines.append(f"{name}: {added} lead(s) added. Skipped: {why}.")
    seq = "They're in the email sequence and get their first email on the next runs." if status == "active" else "They were added without starting the email sequence."
    waiting = con.execute("SELECT COUNT(*) FROM leads WHERE status='active' AND step=0").fetchone()[0]
    answer(f"BSO imported your lead file into the campaign \"{cname}\".\n\n" + "\n".join(lines) + (f"\n\n{seq}" if total else "") +
           f"\n\nLeads waiting for their first email: {waiting}.")
    notify(f"📥 {total} lead(s) imported from a file emailed by {frm}", cfg)
    return True


# ---------------------------------------------------------------- referrals: "talk to my colleague Sarah instead"

def _without_quotes(body):
    out = []
    for line in body.replace("\r\n", "\n").split("\n"):
        if re.match(r"^\s*On .{5,200} wrote:\s*$", line) or line.strip().startswith("-----Original Message-----"):
            break
        if not line.lstrip().startswith(">"):
            out.append(line)
    return "\n".join(out)


def _name_near(text, addr):
    m = re.search(r"([A-Z][a-z'’-]+(?:\s+[A-Z][a-z'’-]+)?)\s*(?:<|\(|\[|:|,|-|–|at|on)?\s*(?:mailto:)?" + re.escape(addr), text)
    if m and m.group(1).lower() not in {"email", "contact", "reach", "at", "to", "on", "write", "ping"}:
        return m.group(1).split()
    local = addr.split("@")[0]
    parts = [p for p in re.split(r"[._-]+", local) if p.isalpha() and len(p) > 1]
    return [p.capitalize() for p in parts[:2]] if parts else []


def capture_referral(cfg, con, lead, sender, body):
    """Add the colleague(s) a prospect pointed us to as new leads. Returns the new lead ids."""
    if not (cfg.get("referral") or {}).get("enabled", True):
        return []
    text = _without_quotes(body)
    ours = {str(sv.get("email", "")).lower() for sv in cfg.get("senders") or []}
    found = []
    for addr in dict.fromkeys(a.lower().rstrip(".") for a in FIND_EMAILS_RE.findall(text)):
        if addr == lead["email"] or addr in ours or not EMAIL_RE.match(addr) or addr.split("@")[0] in ROLE_INBOXES:
            continue
        found.append(addr)
    referrer = " ".join(p for p in (lead.get("first_name"), json.loads(lead.get("extra") or "{}").get("last_name")) if p) or lead["email"]
    new = []
    for addr in found[:2]:
        name = _name_near(text, addr)
        row = {"email": addr, "first_name": name[0] if name else "", "last_name": " ".join(name[1:]) if len(name) > 1 else "",
               "company": lead.get("company") or "", "website": lead.get("website") or "",
               "source": f"referral from {lead['email']}", "referred_by": referrer, "referred_by_email": lead["email"],
               "campaign_id": lead.get("campaign_id") or 1}
        lid = add_lead(con, row, "active")
        if lid:
            new.append(lid)
            con.execute("INSERT INTO notes(lead_id, body, at, kind) VALUES (?,?,?,'note')",
                        (lead["id"], f"Referred us to {' '.join(name) or addr} <{addr}>: added as a new lead.", iso(now_utc())))
            con.execute("INSERT INTO notes(lead_id, body, at, kind) VALUES (?,?,?,'note')",
                        (lid, f"Referred by {referrer} <{lead['email']}>. Their reply:\n{text.strip()[:1500]}", iso(now_utc())))
            log.info("Referral: %s suggested %s", lead["email"], addr)
    con.commit()
    return new


from convo import TEST_TAG_RE  # noqa: E402  (kept importable from here)


def handle_message(cfg, con, sender, msg, folder="INBOX"):
    """Work out whose conversation an incoming email belongs to (by thread first, not by sender) and act on it.
    Returns a short description of what was done (kept in the inbox activity log)."""
    import convo
    inbox = sender["email"]
    frm = parseaddr(msg.get("From", ""))[1].lower()
    body = get_text_body(msg)

    def done(decision, lead_id=None):
        convo.log_inbox(con, inbox, msg, decision, lead_id, folder)
        return decision

    if msg.get("X-BSO-Kind"):
        return done("ignored: an email BSO itself sent (briefing, forward or answer)")
    if "mailer-daemon" in frm or "postmaster" in frm:  # bounce
        hit = []
        for addr in set(a.lower() for a in FIND_EMAILS_RE.findall(body)):
            lead = con.execute("SELECT id FROM leads WHERE email=?", (addr,)).fetchone()
            if lead:
                con.execute("UPDATE leads SET status='bounced' WHERE id=?", (lead["id"],))
                set_stage(con, lead["id"], "unqualified")
                cancel_drafts(con, lead["id"])
                suppress(con, addr, "bounced")
                log_event(con, lead["id"], "bounce", "", sender["email"])
                log.info("Bounce: %s", addr)
                hit.append(addr)
        return done("bounce: " + (", ".join(hit) if hit else "no lead found in it"))

    m = convo.match(cfg, con, msg, frm, inbox, body)
    lead = m["lead"]
    if not lead:
        if folder != "INBOX":
            return done(f"left in {folder}: not part of a BSO conversation")
        if intake_message(cfg, con, sender, msg):
            return done("lead file email (imported or answered)")
        return done("ignored: not from a lead or part of a BSO conversation")
    if folder != "INBOX" and not m["sure"]:
        return done(f"left in {folder}: only loosely matches {lead['email']}", lead["id"])
    lead = dict(lead)
    where = f" (found in {folder})" if folder != "INBOX" else ""
    convo.link_people(con, lead, msg, frm, m["how"], m["sure"], cfg)
    if m["path"]:
        convo.add_path_contacts(con, lead, m["path"])
    who_sent = "" if frm == lead["email"] else f" from {frm}"
    if lead.get("handed_to") and lead["status"] not in ("unsubscribed", "bounced"):
        convo.record_received(con, msg, lead["id"], inbox, folder, frm, body, how=m["how"], path=m["path"])
        forward_to_supervisor(cfg, con, lead, sender, msg, body)   # nothing slips through after a handover
        return done(f"handed-over prospect {lead['email']} wrote again{who_sent}: forwarded to their supervisor{where}", lead["id"])
    if lead["status"] in ("unsubscribed", "bounced"):
        convo.record_received(con, msg, lead["id"], inbox, folder, frm, body, how=m["how"], path=m["path"])
        return done(f"ignored: {lead['email']} is on do-not-contact", lead["id"])

    a = analyze_reply(cfg, msg, body)
    label = a["label"]
    x = json.loads(lead.get("extra") or "{}")
    engaged = con.execute("SELECT COUNT(*) FROM received WHERE lead_id=?", (lead["id"],)).fetchone()[0] + \
        (1 if frm != lead["email"] else 0) + (1 if m["path"] else 0)
    pbk = convo.playbook(cfg)
    temp, score, reason, when = convo.temperature(label, body, lead, x.get("title", ""), engaged, a if a.get("temperature") else None,
                                                  int(pbk["cold"].get("check_back_days") or 90))
    pointer = re.search(r"\b(reach out to|contact|speak (to|with)|talk to|get in touch with|right person|better person|best person|handles|responsible for|in charge of|"
                        r"forward(ed|ing)? (this|you)|loop(ing|ed)? in|cc'?d|copying|introduc)", body.lower())
    if label == "referral" or (label in ("not_interested", "other") and pointer):   # "not me, talk to sarah@..." often reads as not interested
        if capture_referral(cfg, con, lead, sender, body) and label != "referral":
            label = "referral"
            temp = "warm" if temp == "cold" else temp
    via = ""
    if frm != lead["email"]:
        via = f"replied by {parseaddr(msg.get('From', ''))[0] or frm} <{frm}>" + (f", reached via {' → '.join(m['path'])}" if m["path"] else "") + \
              f" (matched by {m['how']})"
    log_event(con, lead["id"], "reply", label, sender["email"])
    convo.record_received(con, msg, lead["id"], inbox, folder, frm, body, label, temp, m["how"], m["path"])
    last = con.execute("SELECT at FROM sent_log WHERE lead_id=? AND test=0 ORDER BY id DESC LIMIT 1", (lead["id"],)).fetchone()
    if last and last["at"] < iso(now_utc() - timedelta(days=60)):
        con.execute("UPDATE leads SET came_back=? WHERE id=?", (iso(now_utc())[:10], lead["id"]))
        j = convo.journey(con, lead["id"])
        f = j["first"] or {}
        con.execute("INSERT INTO notes(lead_id, body, at, kind) VALUES (?,?,?,'note')",
                    (lead["id"], f"Came back {(now_utc().date() - date.fromisoformat(last['at'][:10])).days} days after our last email. "
                     f"First reached {str(f.get('at', ''))[:10]} by campaign \"{f.get('campaign', '')}\" (email {int(f.get('step') or 0) + 1}"
                     f"{', version ' + f['variant'] if f.get('variant') else ''}, subject \"{f.get('subject', '')}\") from {f.get('sender', '')}.", iso(now_utc())))
    if not m["sure"]:
        add_task(con, f"Confirm: {frm} ({m['how']}) wrote to {inbox}. Is it part of {lead['first_name'] or lead['email']}'s conversation?", lead_id=lead["id"])
    if via:
        con.execute("INSERT INTO notes(lead_id, body, at, kind) VALUES (?,?,?,'note')", (lead["id"], via[0].upper() + via[1:] + ".", iso(now_utc())))

    if label == "out_of_office":
        back = convo.return_date(body)
        if back and lead["status"] == "active":
            nxt = datetime.combine(back + timedelta(days=2), datetime.min.time(), tzinfo=timezone.utc).replace(hour=9)
            con.execute("UPDATE leads SET next_send_at=? WHERE id=? AND (next_send_at IS NULL OR next_send_at<?)", (iso(nxt), lead["id"], iso(nxt)))
            return done(f"out of office until {back.isoformat()}: emails to {lead['email']} paused until {(back + timedelta(days=2)).isoformat()}{where}", lead["id"])
        return done(f"out of office: sequence for {lead['email']} carries on{where}", lead["id"])

    cancel_drafts(con, lead["id"])
    who = lead["first_name"] or lead["email"]
    stopped = convo.company_stop(con, lead, f"{who} at {lead['company'] or convo.company_domain(lead['email'])} replied")
    convo.set_temperature(con, lead["id"], temp, score, reason, "bot", when)
    log_event(con, lead["id"], "temperature", temp)
    extra_note = f"; {stopped} other lead(s) at the company paused" if stopped else ""
    if temp == "closed":
        con.execute("UPDATE leads SET status='unsubscribed' WHERE id=?", (lead["id"],))
        set_stage(con, lead["id"], "unqualified")
        suppress(con, lead["email"], label if label in ("unsubscribe", "not_interested") else "asked not to be contacted")
        if frm != lead["email"]:
            suppress(con, frm, "asked not to be contacted")
        return done(f"{label.replace('_', ' ')} from {lead['email']}{who_sent}: closed, on do-not-contact{extra_note}{where}", lead["id"])
    if temp == "cold":
        con.execute("UPDATE leads SET status='nurture' WHERE id=?", (lead["id"],))
        set_stage(con, lead["id"], "replied")
        con.execute("INSERT INTO notes(lead_id, body, at, kind) VALUES (?,?,?,'email')",
                    (lead["id"], f"[{label} reply, cold: {reason}]\n{body[:2000]}", iso(now_utc())))
        add_task(con, f"Check back with {who} ({reason})", due_date=when.isoformat(), lead_id=lead["id"])
        con.execute("UPDATE tasks SET kind='nurture' WHERE id=last_insert_rowid()")
        notify(f"❄ Cold reply from {who} ({lead['company'] or ''}): {reason}. Check back on {when.isoformat()}.", cfg)
        return done(f"❄ cold reply from {lead['email']}{who_sent} ({reason}): nurture, check back {when.isoformat()}{extra_note}{where}", lead["id"])

    con.execute("UPDATE leads SET status='replied' WHERE id=?", (lead["id"],))
    set_stage(con, lead["id"], "replied")
    con.execute("INSERT INTO notes(lead_id, body, at, kind) VALUES (?,?,?,'email')",
                (lead["id"], f"[{label} reply, {temp}]\n{body[:2000]}", iso(now_utc())))
    icon = convo.TEMP_ICON[temp]
    if label in ("question", "interested", "other") and temp != "hot":
        try:
            import ops
            answered = ops.info_reply(cfg, con, lead, sender, msg, body)
        except Exception as e:
            log.warning("Information answer failed, handing over as usual: %s", e)
            answered = None
        if answered:
            return done(f"{icon} {answered}{extra_note}{where}", lead["id"])
    if (label in HANDOVER_LABELS or temp == "hot") and handover(cfg, con, lead, sender, msg, body, label, temp, reason):
        return done(f"{icon} {temp} {label.replace('_', ' ')} reply from {lead['email']}{who_sent}: handed over{extra_note}{where}", lead["id"])
    pb = convo.playbook(cfg)[temp]
    due = convo.add_working_hours(datetime.now(), pb["respond_hours"])
    add_task(con, f"{icon} Respond to {who} ({temp}, {label.replace('_', ' ')} reply) within {pb['respond_hours']} working hour(s)",
             due_date=due.date().isoformat(), lead_id=lead["id"])
    con.execute("UPDATE tasks SET priority=?, due_at=?, kind='respond' WHERE id=last_insert_rowid()", ("high" if temp == "hot" else "normal", due.isoformat(timespec="minutes")))
    snippet = body[:300].replace("\n", " ")
    notify(f"📬 {icon} {temp.upper()} {label.upper()} reply from {lead['first_name'] or ''} <{frm}> "
           f"({lead['company'] or ''}) in {sender['email']}:\n{snippet}", cfg)
    return done(f"{icon} {temp} {label.replace('_', ' ')} reply from {lead['email']}{who_sent}: task to respond{extra_note}{where}", lead["id"])


# ---------------------------------------------------------------- handover to a lead gen supervisor

DEFAULT_INTRO = """Hi {first_name},

Thanks for getting back to me. I'm looping in my colleague {supervisor_name}, who works with our clients on exactly this. {supervisor_first_name} will take it from here and can find a time that suits you.

Best,
{sender_name}"""


def supervisors(cfg):
    """[{id, name, email, count}]: from the master console when on a team, else from Settings → Handover."""
    try:
        import team
        if team.joined():
            return team.supervisors()
    except Exception as e:
        log.warning("Team supervisors unavailable: %s", e)
    out = []
    for sv in (cfg.get("handover") or {}).get("supervisors") or []:
        if EMAIL_RE.match(str(sv.get("email", "")).strip()):
            out.append({"id": sv["email"].strip().lower(), "name": sv.get("name") or sv["email"], "email": sv["email"].strip().lower(), "count": 0})
    return out


def pick_supervisor(cfg, con):
    """Take turns: the supervisor with the fewest handovers gets the next one."""
    sups = supervisors(cfg)
    if not sups:
        return None
    since = ""
    try:
        import team
        since = team.state().get("sup_as_of", "") if team.joined() else ""
    except Exception:
        pass
    def load(sv):
        local = con.execute("SELECT COUNT(*) FROM leads WHERE handed_to=? AND handed_at>?", (sv["email"], since)).fetchone()[0]
        return (sv.get("count", 0) + local, sups.index(sv))
    return min(sups, key=load)


def _thread(con, lead_id):
    return [{"at": r["sent_at"] or r["created_at"], "subject": r["subject"], "body": r["body"]}
            for r in con.execute("SELECT * FROM outbox WHERE lead_id=? AND status='sent' ORDER BY id", (lead_id,))]


def handover(cfg, con, lead, sender, msg, body, label, temp="warm", reason=""):
    """Hand an interested prospect to a lead gen supervisor. Returns False when there's nobody to hand to."""
    import convo
    hc = cfg.get("handover") or {}
    if not hc.get("enabled", True):
        return False
    sup = None
    cd = convo.company_domain(lead["email"])
    if cd:                                              # one company stays with one supervisor
        owner = con.execute("SELECT handed_to FROM leads WHERE handed_to IS NOT NULL AND id!=? AND lower(email) LIKE ? ORDER BY handed_at DESC LIMIT 1",
                            (lead["id"], "%@" + cd)).fetchone()
        if owner:
            sup = next((x for x in supervisors(cfg) if x["email"].lower() == owner["handed_to"].lower()), None)
    try:
        if sup:
            raise StopIteration

        import team
        sup = team.assign_supervisor(lead["email"])     # the always-on relay keeps the turns fair for everyone
    except StopIteration:
        pass
    except Exception as e:
        log.warning("Supervisor assignment unavailable: %s", e)
    sup = sup or pick_supervisor(cfg, con)
    if not sup:
        return False
    tm = cfg.get("test_mode") or {}
    redirect = tm.get("redirect_to") if tm.get("enabled") else None
    name = " ".join(p for p in (lead.get("first_name"), json.loads(lead.get("extra") or "{}").get("last_name")) if p) or lead["email"]
    company = lead.get("company") or ""
    sup_first = (sup["name"] or "").split()[0] if sup.get("name") else "they"

    # 1. warm intro from the same inbox, in the same thread, supervisor in CC
    intro_note = ""
    intro_text = ""
    if hc.get("intro_enabled", True):
        tpl = (hc.get("intro_template") or DEFAULT_INTRO).replace("{supervisor_name}", sup["name"]).replace("{supervisor_first_name}", sup_first)
        intro_text = render(tpl, lead, sender, cfg)
        hits = find_avoid(cfg, intro_text)
        subject = msg.get("Subject", "") or ("Re: " + (lead.get("first_subject") or ""))
        subject = subject if subject.lower().startswith("re:") else "Re: " + subject
        try:
            if redirect:
                mid = send_email(sender, redirect, f"[TEST for {lead['email']}, CC {sup['email']}] {subject}", intro_text, msg.get("Message-ID"))
            else:
                mid = send_email(sender, lead["email"], subject, intro_text, msg.get("Message-ID"), cc=sup["email"])
            convo.record_sent(con, mid, lead, "intro", sender["email"], redirect or lead["email"], subject, intro_text,
                              cc=None if redirect else sup["email"], test=bool(redirect))
            log_event(con, lead["id"], "sent", "handover intro" + (" (test)" if redirect else ""), sender["email"])
            if hits:
                log.warning("Handover intro uses avoided words: %s", ", ".join(hits))
        except Exception as e:
            log.error("Handover intro to %s failed: %s", lead["email"], e)
            intro_note = f"\n⚠ The introduction email could not be sent ({e}). Please contact the prospect directly.\n"
            intro_text = ""

    # 2. brief for the supervisor
    thread = _thread(con, lead["id"])
    x = json.loads(lead.get("extra") or "{}")
    member = ""
    try:
        import team
        member = (team.policy() or {}).get("name", "") if team.joined() else ""
    except Exception:
        pass
    lines = [f"{sup['name']}, a prospect has been handed over to you.", intro_note,
             f"Prospect:   {name}", f"Title:      {x.get('title', '')}", f"Company:    {company}", f"Email:      {lead['email']}",
             f"Website:    {lead.get('website') or ''}", f"LinkedIn:   {x.get('linkedin', '')}",
             f"Reply type: {label.replace('_', ' ')}", f"Temperature: {convo.TEMP_ICON.get(temp, '')} {temp} ({reason})",
             f"Respond within: {convo.playbook(cfg)[temp]['respond_hours']} working hour(s)" if temp in ("hot", "warm") else "",
             f"From:       {member or sender['name']} ({sender['email']})", "",
             "--- Their reply ---", body.strip()[:4000], ""]
    if intro_text:
        lines += ["--- Introduction we sent (you were in CC) ---", intro_text, ""]
    for i, t in enumerate(thread, 1):
        lines += [f"--- Our email {i} ({str(t['at'])[:10]}): {t['subject']} ---", t["body"], ""]
    lines.append("Reply to the prospect from your own mailbox. In BSO, open My prospects to track the stage, notes and tasks.")
    brief_to = redirect or sup["email"]
    try:
        send_email(sender, brief_to, f"{'[TEST] ' if redirect else ''}Handover: {name}, {company} ({label.replace('_', ' ')})", "\n".join(lines),
                   headers={"X-BSO-Kind": "brief", "X-BSO-Prospect": lead["email"]})
    except Exception as e:
        log.error("Handover brief to %s failed: %s", sup["email"], e)

    # 3. record it
    at = iso(now_utc())
    con.execute("UPDATE leads SET handed_to=?, handed_at=? WHERE id=?", (sup["email"], at, lead["id"]))
    set_stage(con, lead["id"], "qualified", auto=False)
    con.execute("INSERT INTO notes(lead_id, body, at, kind) VALUES (?,?,?,'note')",
                (lead["id"], f"Handed over to {sup['name']} <{sup['email']}> ({label.replace('_', ' ')} reply).", at))
    log_event(con, lead["id"], "handover", sup["id"], sender["email"])
    j = convo.journey(con, lead["id"])
    first = j["first"] or {}
    suggestion = ""
    try:
        import writer
        suggestion = writer.suggest_reply(cfg, lead, body, temp)
    except Exception as e:
        log.warning("Suggested reply unavailable: %s", e)
    payload = {"lead_local_id": lead["id"], "supervisor": sup["id"], "label": label, "at": at, "temperature": temp, "reason": reason,
               "respond_hours": convo.playbook(cfg).get(temp, {}).get("respond_hours"), "suggested_reply": suggestion,
               "first_touch": {"campaign": first.get("campaign"), "at": first.get("at"), "step": first.get("step"), "variant": first.get("variant"),
                               "subject": first.get("subject"), "sender": first.get("sender")},
               "related": [dict(r) for r in con.execute("SELECT email, name, role FROM related_contacts WHERE lead_id=?", (lead["id"],))],
               "campaign": campaign_row(con, lead.get("campaign_id"))["name"],
               "lead": {"email": lead["email"], "first_name": lead.get("first_name"), "last_name": x.get("last_name"), "title": x.get("title"),
                        "company": company, "website": lead.get("website"), "linkedin": x.get("linkedin"), "city": x.get("city"), "country": x.get("country")},
               "reply": body.strip()[:4000], "intro": intro_text, "thread": thread, "from_inbox": sender["email"], "from_name": sender["name"]}
    con.execute("INSERT INTO handover_out(lead_id, supervisor, payload, at, synced) VALUES (?,?,?,?,?)",
                (lead["id"], sup["id"], json.dumps(payload), at, 1 if redirect else 0))   # a test handover stays on this computer
    con.commit()
    notify(f"🤝 Handed over {name} ({company}) to {sup['name']}: {label.replace('_', ' ')} reply in {sender['email']}", cfg)
    return True


def forward_to_supervisor(cfg, con, lead, sender, msg, body):
    """A handed-over prospect wrote to the original inbox again: pass it to their supervisor straight away."""
    tm = cfg.get("test_mode") or {}
    redirect = tm.get("redirect_to") if tm.get("enabled") else None
    name = lead.get("first_name") or lead["email"]
    text = (f"{name} <{lead['email']}> wrote to {sender['email']} after being handed over to you.\n"
            f"Reply to them directly (Reply goes to {lead['email']}).\n\n--- Their message ---\n{body.strip()[:8000]}")
    try:
        send_email(sender, redirect or lead["handed_to"], f"{'[TEST] ' if redirect else ''}Fwd: {msg.get('Subject', '')} ({name} replied to {sender['email']})",
                   text, headers={"X-BSO-Kind": "forward", "X-BSO-Prospect": lead["email"], "Reply-To": lead["email"]})
        log_event(con, lead["id"], "reply", "after handover: forwarded to supervisor", sender["email"])
    except Exception as e:
        log.error("Forward to supervisor %s failed: %s", lead["handed_to"], e)
        add_task(con, f"Pass on a new email from {name} to {lead['handed_to']} (forwarding failed)", lead_id=lead["id"])
    if classify_reply(cfg, msg, body) == "unsubscribe":
        suppress(con, lead["email"], "unsubscribe")
    con.commit()


JUNK_NAMES = ("spam", "junk", "bulk", "junk e-mail", "junk email")


def _junk_folder(M):
    """The Spam/Junk folder's name on this server, or None."""
    try:
        _, rows = M.list()
    except Exception:
        return None
    best = None
    for raw in rows or []:
        line = raw.decode(errors="replace") if isinstance(raw, bytes) else str(raw)
        m = re.match(r'^\((?P<flags>[^)]*)\)\s+(?:"[^"]*"|NIL)\s+(?P<name>.+)$', line)
        if not m:
            continue
        name = m.group("name").strip()
        bare = name.strip('"')
        if "\\Junk" in m.group("flags"):
            return name
        if bare.lower().split(".")[-1].split("/")[-1] in JUNK_NAMES:
            best = best or name
    return best


def process_inbox(cfg, con, sender, recheck=False, days=None):
    """Read new messages in one inbox (and its Spam folder, for replies to BSO's own threads). Returns how many were new.
    recheck=True looks again at everything from the last `days` days that hasn't been handled as a reply yet."""
    n_new = 0
    since = (datetime.now() - timedelta(days=days or cfg.get("reply_lookback_days", 14))).strftime("%d-%b-%Y")
    M = imaplib.IMAP4_SSL(sender["imap_host"], sender.get("imap_port", 993))
    try:
        M.login(sender.get("username", sender["email"]), sender["password"])
        folders = ["INBOX"]
        junk = _junk_folder(M)
        if junk and junk.strip('"').upper() != "INBOX":
            folders.append(junk)
        for folder in folders:
            try:
                typ, _ = M.select(folder, readonly=True)  # never marks your mail as read
            except Exception:
                continue
            if typ != "OK":
                continue
            label = "INBOX" if folder == "INBOX" else "Spam"
            _, data = M.uid("search", None, f'(SINCE "{since}")')
            for uid in (data[0].split() if data and data[0] else []):
                # cheap header fetch first; full body only for new messages
                _, hdr = M.uid("fetch", uid, "(BODY.PEEK[HEADER.FIELDS (MESSAGE-ID SUBJECT)])")
                head = email.message_from_bytes(hdr[0][1]) if hdr and hdr[0] and isinstance(hdr[0], tuple) else {}
                mid = (head.get("Message-ID") or "").strip()
                key = mid or f"{sender['email']}:{label}:{uid.decode()}"
                if "[TEST for" in str(head.get("Subject") or ""):
                    key = "t:" + key                 # test replies are matched differently since v1.0.45: look at them once more
                # claim it first, so the live reply watcher and a run never handle the same message twice
                claimed = con.execute("INSERT OR IGNORE INTO seen_messages VALUES (?)", (key,)).rowcount
                con.commit()                         # always end the transaction, or other readers of this inbox would wait
                if not claimed:
                    if not recheck or not mid:
                        continue
                    done_before = con.execute("SELECT 1 FROM received WHERE msg_id=?", (mid,)).fetchone() or con.execute(
                        "SELECT 1 FROM inbox_log WHERE msg_id=? AND (decision LIKE 'bounce%' OR decision LIKE 'lead file%' OR decision LIKE '%reply%' "
                        "OR decision LIKE 'out of office%' OR decision LIKE 'handed-over%')", (mid,)).fetchone()
                    if done_before:
                        continue
                _, full = M.uid("fetch", uid, "(BODY.PEEK[])")
                msg = email.message_from_bytes(full[0][1])
                try:
                    handle_message(cfg, con, sender, msg, label)
                    con.commit()
                except Exception as e:
                    con.rollback()
                    log.error("Couldn't handle a message in %s: %s", sender["email"], e)
                    import convo
                    convo.log_inbox(con, sender["email"], msg, f"error: {e}"[:200], None, label)
                    con.commit()
                n_new += 1
    finally:
        try:
            M.logout()
        except Exception:
            pass
    return n_new


def recheck_once(cfg, con, sender):
    """After updating to this version, look once more at the last 14 days of each inbox (replies an older version skipped)."""
    k = f"rechecked_v2:{sender['email'].lower()}"
    if con.execute("SELECT 1 FROM app_mode WHERE k=?", (k,)).fetchone():
        return 0
    n = process_inbox(cfg, con, sender, recheck=True)
    con.execute("INSERT OR REPLACE INTO app_mode(k, v) VALUES (?, '1')", (k,))
    con.commit()
    return n


def cmd_check(cfg, con, args):
    """Read every inbox for replies. With many inboxes (the BSO Bot) several are read at the same time."""
    senders = [s for s in load_senders(cfg) if s["password"]]
    def one(s, c):
        try:
            recheck_once(cfg, c, s)
            process_inbox(cfg, c, s)
        except Exception as e:
            log.error("Inbox check failed for %s: %s", s["email"], e)
    workers = int(((cfg.get("autopilot") or {}).get("check_workers")) or 8)
    if len(senders) <= 3 or workers <= 1:
        for s in senders:
            one(s, con)
        con.commit()
        return
    con.commit()
    from concurrent.futures import ThreadPoolExecutor
    def run(s):
        c = db_connect(cfg.get("database", "outreach.db"))
        try:
            one(s, c)
            c.commit()
        finally:
            c.close()
    with ThreadPoolExecutor(max_workers=workers) as ex:
        list(ex.map(run, senders))


# ---------------------------------------------------------------- other commands

HEADER_ALIASES = {
    "email": "email", "email address": "email", "work email": "email", "e-mail": "email",
    "first name": "first_name", "firstname": "first_name", "first_name": "first_name",
    "last name": "last_name", "lastname": "last_name", "last_name": "last_name",
    "company": "company", "company name": "company", "organization": "company", "account": "company",
    "title": "title", "job title": "title", "position": "title",
    "website": "website", "company website": "website", "domain": "website", "url": "website",
    "employee linkedin": "linkedin", "linkedin": "linkedin", "person linkedin": "linkedin",
    "linkedin url": "linkedin", "company li": "company_linkedin", "company linkedin": "company_linkedin",
    "industry": "industry", "city": "city", "country": "country", "source": "source",
    # Instantly exports (SuperSearch / lead lists / campaign leads)
    "companyname": "company", "jobtitle": "title", "company domain": "website", "companydomain": "website",
    "linkedin profile": "linkedin", "linkedinurl": "linkedin", "location": "city", "company_name": "company",
    "job_title": "title", "website url": "website",
    "campaign_code": "campaign_code", "campaign code": "campaign_code", "campaign": "campaign_code",
    "company_size": "company_size", "company size": "company_size", "employees": "company_size",
    "linkedin_url": "linkedin", "notes": "notes", "time zone": "timezone", "timezone": "timezone",
}
ROLE_INBOXES = {"info", "hello", "help", "support", "sales", "contact", "admin", "office",
                "team", "hr", "careers", "jobs", "marketing", "billing", "noreply", "no-reply"}
FIT_RULES = [  # (grade, title keywords) — first match wins
    ("A", ["cto", "chief technology", "vp engineering", "vp of engineering", "head of engineering",
           "director of engineering", "head of product", "cpo", "chief product", "vp technology",
           "head of it", "cio", "chief information", "technology architect", "engineering"]),
    ("B", ["ceo", "chief executive", "founder", "co-founder", "president", "coo", "owner",
           "managing director", "executive director", "partner", "technology", "product"]),
]


def fit_grade(title):
    """A = technical decision maker, B = executive/founder, C = outside target roles.
    Whole-word matching, so 'Director' doesn't count as 'CTO'."""
    t = (title or "").lower()
    for grade, words in FIT_RULES:
        if any(re.search(r"\b" + re.escape(w) + r"\b", t) for w in words):
            return grade
    return "C"


def clean_website(url, email=""):
    """'http://www.Foo.com/about' -> 'foo.com'. Falls back to the email's domain if broken."""
    host = re.sub(r"^https?://", "", (url or "").strip().lower()).split("/")[0].removeprefix("www.")
    if not re.match(r"^[a-z0-9.-]+\.[a-z]{2,}$", host):
        host = email.split("@")[-1] if "@" in email else ""
    return host


def normalize_row(row):
    """Map any reasonable spreadsheet headers to our fields and tidy the values."""
    out = {}
    for k, v in row.items():
        if k is None:
            continue
        key = str(k).strip().lower()
        key = HEADER_ALIASES.get(key, key.replace(" ", "_"))
        out[key] = "" if v is None else str(v).strip()
    for f in ("first_name", "last_name"):
        if out.get(f, "").isupper() and len(out[f]) > 2:
            out[f] = out[f].title()  # ALLEN -> Allen
    out["email"] = out.get("email", "").lower()
    out["website"] = clean_website(out.get("website", ""), out["email"])
    return out


def add_lead(con, row, status="active", why=None):
    """Insert one lead from a dict. Returns the new id, or None if invalid/duplicate/suppressed.
    If `why` is a list, the skip reason is appended to it."""
    row = normalize_row(row)
    addr = row.get("email", "")
    reason = None
    if not EMAIL_RE.match(addr):
        reason = "invalid email"
    elif addr.split("@")[0] in ROLE_INBOXES:
        reason = "shared inbox (e.g. info@, hello@)"
    elif con.execute("SELECT 1 FROM suppression WHERE email=?", (addr,)).fetchone():
        reason = "on do-not-contact list"
    elif con.execute("SELECT 1 FROM leads WHERE email=?", (addr,)).fetchone():
        reason = "already in CRM"
    if reason:
        if why is not None:
            why.append(reason)
        return None
    row["fit"] = fit_grade(row.get("title"))
    try:
        campaign = int(row.pop("campaign_id", None) or 1)
    except (TypeError, ValueError):
        campaign = 1
    extra = {k: v for k, v in row.items() if k not in CORE_FIELDS and v}
    cur = con.execute(
        "INSERT OR IGNORE INTO leads(email, first_name, company, website, extra, status, "
        "next_send_at, created_at, campaign_id) VALUES (?,?,?,?,?,?,?,?,?)",
        (addr, row.get("first_name"), row.get("company"), row.get("website"),
         json.dumps(extra), status, iso(now_utc()), iso(now_utc()), campaign),
    )
    if not cur.rowcount:
        return None
    lid = cur.lastrowid
    try:
        import convo
        prev = convo.previously(con, addr)
        if prev["touches"] and status == "active":
            t0, t1 = prev["touches"][0], prev["touches"][-1]
            camp = con.execute("SELECT name FROM campaigns WHERE id=?", (t0.get("campaign_id") or 1,)).fetchone()
            note = (f"Already contacted before: {len(prev['touches'])} email(s), first on {t0['at'][:10]} "
                    f"(campaign \"{camp['name'] if camp else ''}\", subject \"{t0['subject']}\"), last on {t1['at'][:10]}. "
                    "Not starting the sequence again: write a reconnect email from the lead's page instead.")
            con.execute("UPDATE leads SET status='paused', extra=json_set(extra, '$.previously_contacted', ?) WHERE id=?", (t1["at"][:10], lid))
            con.execute("INSERT INTO notes(lead_id, body, at, kind) VALUES (?,?,?,'note')", (lid, note, iso(now_utc())))
            add_task(con, f"Previously contacted: {addr}. Reconnect instead of a cold first email?", lead_id=lid)
            con.execute("UPDATE tasks SET kind='reconnect_offer' WHERE id=last_insert_rowid()")
        elif prev["company_contacts"]:
            con.execute("UPDATE leads SET extra=json_set(extra, '$.company_contacted_before', ?) WHERE id=?", (prev["company_contacts"], lid))
    except Exception as e:
        log.warning("Touch history check failed for %s: %s", addr, e)
    return lid


def read_rows(path_or_file, filename=""):
    """Yield dict rows from a CSV or Excel file (path or file-like object)."""
    name = (filename or str(path_or_file)).lower()
    if name.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(path_or_file, read_only=True, data_only=True)
        ws = next((w for w in wb.worksheets if w.title.strip().lower() == "leads"), wb.active)   # BSO's template: the Leads sheet
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None) or []
        for r in rows:
            if any(c not in (None, "") for c in r):
                yield dict(zip(header, r))
    elif name.endswith(".xls"):
        import xlrd                      # older Excel files
        data = path_or_file.read() if hasattr(path_or_file, "read") else Path(path_or_file).read_bytes()
        book = xlrd.open_workbook(file_contents=data)
        sh = next((book.sheet_by_name(n) for n in book.sheet_names() if n.strip().lower() == "leads"), book.sheet_by_index(0))
        header = [str(c).strip() for c in sh.row_values(0)] if sh.nrows else []
        for i in range(1, sh.nrows):
            r = sh.row_values(i)
            if any(c not in (None, "") for c in r):
                yield dict(zip(header, [str(int(c)) if isinstance(c, float) and c.is_integer() else c for c in r]))
    else:
        if isinstance(path_or_file, (str, Path)):
            with open(path_or_file, newline="", encoding="utf-8-sig") as f:
                yield from csv.DictReader(f)
        else:
            yield from csv.DictReader(path_or_file)


def import_rows(con, rows, status="active", source="", campaign_id=1):
    """Returns (added, {reason: count})."""
    added, reasons = 0, []
    for row in rows:
        if source and not any(str(k).strip().lower() == "source" for k in row if k):
            row = {**row, "source": source}
        row = {**row, "campaign_id": campaign_id}
        if add_lead(con, row, status, reasons):
            added += 1
    con.commit()
    summary = {}
    for r in reasons:
        summary[r] = summary.get(r, 0) + 1
    return added, summary


def import_csv(con, f, status="active"):  # kept for older callers
    added, skipped = import_rows(con, csv.DictReader(f), status)
    return added, sum(skipped.values())


def cmd_import(cfg, con, args):
    added, skipped = import_rows(con, read_rows(args.file), source=Path(args.file).name)
    print(f"Imported {added} lead(s).")
    for reason, n in skipped.items():
        print(f"  skipped {n}: {reason}")


def cmd_suppress(cfg, con, args):
    suppress(con, args.email, "manual")
    con.execute("UPDATE leads SET status='unsubscribed' WHERE email=?", (args.email.lower(),))
    row = con.execute("SELECT id FROM leads WHERE email=?", (args.email.lower(),)).fetchone()
    if row:
        set_stage(con, row["id"], "unqualified")
        cancel_drafts(con, row["id"])
    con.commit()
    print(f"{args.email} will never be emailed.")


def cmd_selfcheck(cfg, con, args):
    """Verify every library the app loads on demand is present (used by the Windows build test)."""
    problems = []
    mods = ["anthropic", "bs4", "openpyxl", "yaml", "requests", "dotenv", "flask", "apollo", "legal", "secretstore", "pystray", "PIL",
            "autopilot", "ops", "updater", "crm_roles", "hub_bot"]
    if getattr(sys, "frozen", False):
        mods += ["xlrd", "pypdf", "holidays"]          # older Excel files, scope PDFs, public holidays by country
    for mod in mods:
        try:
            __import__(mod)
        except Exception as e:
            problems.append(f"{mod}: {e}")
    try:
        import autopilot
        import updater
        assert updater._key(), "update_key.json missing"
        assert not autopilot.working_day("GB", date(2026, 12, 25))
        if getattr(sys, "frozen", False):
            import holidays
            assert date(2026, 12, 25) in holidays.country_holidays("GB", years=2026)
    except Exception as e:
        problems.append(f"bot rules / updates: {e}")
    try:
        ZoneInfo(cfg["schedule"]["timezone"])
        in_send_window(cfg)
    except Exception as e:
        problems.append(f"time zones: {e}")
    try:  # secrets round-trip through the encrypted store
        import secretstore
        st = secretstore.SecretStore(DATA)
        before = st.load()
        st.set("SELFCHECK_PROBE", "probe-1234")
        assert secretstore.SecretStore(DATA).load().get("SELFCHECK_PROBE") == "probe-1234"
        assert b"probe-1234" not in (DATA / "secrets.dat").read_bytes() or not secretstore.ENCRYPTED
        st.save(before)
        print("secrets: encrypted with Windows DPAPI" if secretstore.ENCRYPTED else "secrets: development mode (not encrypted)")
    except Exception as e:
        problems.append(f"secret store: {e}")
    try:
        import legal
        for d in legal.DOCS:
            assert "{{" not in legal.source(d, cfg)
        print("legal documents: OK, version", legal.app_version())
    except Exception as e:
        problems.append(f"legal documents: {e}")
    try:
        bad = template_avoid_hits(cfg)
        assert not bad, "; ".join(f"step {i} '{v}': {', '.join(h)}" for i, vs in bad.items() for v, h in vs.items())
        if cfg.get("avoid_phrases"):
            assert find_avoid(cfg, "Act now, it is FREE!"), "avoid-list matcher found nothing in a test sentence"
        print("email templates: no avoided phrases (%d phrases checked)" % len(cfg.get("avoid_phrases") or []))
    except Exception as e:
        problems.append(f"email templates use avoided phrases: {e}")
    for p in problems:
        print("PROBLEM:", p)
    print("SELFCHECK OK" if not problems else "SELFCHECK FAILED")
    if problems:
        sys.exit(1)


def cmd_apollo_preview(cfg, con, args):
    import apollo
    try:
        apollo.preview(cfg)
    except apollo.ApolloError as e:
        sys.exit(f"Apollo: {e}")


def cmd_apollo_pull(cfg, con, args):
    import apollo
    try:
        apollo.pull(cfg, con, args.max or cfg["apollo"].get("max_per_pull", 25), add_lead)
    except apollo.ApolloError as e:
        sys.exit(f"Apollo: {e}")


def cmd_stats(cfg, con, args):
    print("Leads by status:")
    for r in con.execute("SELECT status, COUNT(*) n FROM leads GROUP BY status ORDER BY n DESC"):
        print(f"  {r['status']:<15} {r['n']}")
    print("\nEmails sent by step:")
    for r in con.execute("SELECT detail, COUNT(*) n FROM events WHERE type='sent' GROUP BY detail ORDER BY detail"):
        print(f"  {r['detail']:<15} {r['n']}")
    print("\nReplies by type:")
    for r in con.execute("SELECT detail, COUNT(*) n FROM events WHERE type='reply' GROUP BY detail ORDER BY n DESC"):
        print(f"  {r['detail']:<15} {r['n']}")
    contacted = con.execute("SELECT COUNT(DISTINCT lead_id) FROM events WHERE type='sent'").fetchone()[0]
    replied = con.execute(
        "SELECT COUNT(DISTINCT lead_id) FROM events WHERE type='reply' AND detail!='out_of_office'"
    ).fetchone()[0]
    bounced = con.execute("SELECT COUNT(*) FROM leads WHERE status='bounced'").fetchone()[0]
    if contacted:
        print(f"\nContacted {contacted} | reply rate {replied / contacted:.1%} | bounce rate {bounced / contacted:.1%}")


# ---------------------------------------------------------------- main

def acquire_lock():
    """Stop two overlapping cron runs from double-sending."""
    try:
        import fcntl
    except ImportError:  # Windows: no lock
        return None
    fh = open(DATA / ".bot.lock", "w")
    try:
        fcntl.flock(fh, fcntl.LOCK_EX | fcntl.LOCK_NB)
    except OSError:
        print("Another run is in progress; exiting.")
        sys.exit(0)
    return fh


def main(argv=None):
    job_log = os.getenv("OUTREACH_JOB_LOG")
    if job_log:
        sys.stdout = sys.stderr = open(job_log, "a", buffering=1, encoding="utf-8")
    p = argparse.ArgumentParser(description="Cold email outreach bot")
    p.add_argument("--config", default=str(DATA / "config.yaml"))
    p.add_argument("-v", "--verbose", action="store_true")
    sub = p.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("import", help="import leads from a CSV or Excel file")
    s.add_argument("file")
    s = sub.add_parser("run", help="check replies, then send due emails")
    s.add_argument("--dry-run", action="store_true", help="print emails instead of sending")
    s.add_argument("--force", action="store_true", help="ignore the sending window")
    s.add_argument("--no-check", action="store_true", help="skip the reply check")
    s.add_argument("--limit", type=int, default=50, help="max leads to process this run")
    sub.add_parser("check", help="check inboxes for replies and bounces")
    sub.add_parser("stats", help="show campaign stats")
    s = sub.add_parser("suppress", help="never email an address")
    s.add_argument("email")
    sub.add_parser("selfcheck", help="verify the installation")
    sub.add_parser("apollo-preview", help="free: see what your Apollo searches match")
    s = sub.add_parser("apollo-pull", help="add new verified leads from Apollo (uses credits)")
    s.add_argument("--max", type=int, help="max new leads this run (default: apollo.max_per_pull)")

    args = p.parse_args(argv)
    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO,
                        format="%(asctime)s %(levelname)s %(message)s")
    cfg = load_config(args.config)
    con = db_connect(cfg.get("database", "outreach.db"))
    sync_mode(con, cfg)
    lock = acquire_lock() if args.cmd in ("run", "check") else None

    {"import": cmd_import, "run": cmd_run, "check": cmd_check,
     "stats": cmd_stats, "suppress": cmd_suppress,
     "apollo-preview": cmd_apollo_preview, "selfcheck": cmd_selfcheck, "apollo-pull": cmd_apollo_pull}[args.cmd](cfg, con, args)


if __name__ == "__main__":
    main()
