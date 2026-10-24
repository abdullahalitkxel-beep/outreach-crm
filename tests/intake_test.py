"""Lead intake mailbox (owner only, verified sender, template, row checks, receipt with rejected rows, STOP/START ALL),
the 10-day low-stock alert with the template attached, the daily digest, inbox readiness and information answers.
   python tests/intake_test.py"""
import csv, io, json, os, sys, tempfile
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from pathlib import Path

root = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(root))
os.environ["OUTREACH_HOME"] = tempfile.mkdtemp(prefix="intake-test-")
os.environ["BSO_NO_SIGNIN"] = "1"
import bot, crm, autopilot, ops  # noqa: E402,E401

n = 0
def gh_error(text):
    if os.getenv("GITHUB_ACTIONS"):
        print("::error title=intake_test::" + text.replace("%", "%25").replace("\r", "").replace("\n", "%0A")[-3000:], flush=True)
def check(cond, what):
    global n
    print(("PASS " if cond else "FAIL ") + what, flush=True)
    if not cond:
        gh_error("FAIL " + what)
        sys.exit(1)
    n += 1
def _crash(t, v, tb):
    import traceback
    gh_error("".join(traceback.format_exception(t, v, tb)))
    sys.__excepthook__(t, v, tb)
sys.excepthook = _crash

inb = lambda e: {"email": e, "name": "Sam", "smtp_host": "h", "smtp_port": 465, "imap_host": "h", "password_env": "PW1", "daily_limit": 10}
crm.write_settings({"legal_acceptance": {"version": str(crm.cfg().get("legal", {}).get("version", "1.0"))},
                    "senders": [inb("sam@neo-mail.com")], "test_mode": {"enabled": False}, "review": {"enabled": False},
                    "autopilot": {"enabled": True},
                    "intake_box": {"email": "leads@blacswanken.com", "name": "BSO Bot", "imap_host": "imap.x", "smtp_host": "smtp.x",
                                   "smtp_port": 465, "password_env": "INTAKE_PASSWORD", "owner": "omar@blacswanken.com"}})
bot.SECRETS.set("PW1", "pw")
bot.SECRETS.set("INTAKE_PASSWORD", "pw")
os.environ["INTAKE_PASSWORD"] = "pw"
os.environ["PW1"] = "pw"
cfg = lambda: bot.load_config(crm.CONFIG_PATH)
con = bot.db_connect("outreach.db")
smtp_out = []
ops._smtp_send = lambda b, msg: smtp_out.append(msg)
bot.notify = lambda *a, **k: None
ops.has_mail_server = lambda d: d != "nomail.example"

cl = autopilot.add_client(con, "NeoSOFT", "tech services")
autopilot.assign_domain(con, "neo-mail.com", cl)
con.execute("UPDATE campaigns SET name='Neo UK', code='NEO-SMB-UK', client_id=? WHERE id=1", (cl,))
c2 = con.execute("INSERT INTO campaigns(name, status, client_id, created_at, code) VALUES ('Neo US', 'active', ?, '2026-01-01', 'NEO-US')", (cl,)).lastrowid
autopilot.set_campaign_domains(con, 1, ["neo-mail.com"])
autopilot.set_campaign_domains(con, c2, ["neo-mail.com"])
con.commit()

# ---------------------------------------------------------------- the template
import openpyxl
wb = openpyxl.load_workbook(io.BytesIO(ops.template_xlsx(con, cfg())))
check(wb.sheetnames == ["Leads", "Instructions", "Campaigns"], "the template has Leads, Instructions and Campaigns sheets")
codes = [r[0] for r in wb["Campaigns"].iter_rows(min_row=2, values_only=True)]
heads = [c.value for c in wb["Leads"][1]]
check(codes == ["NEO-SMB-UK", "NEO-US"] and heads[:3] == ["campaign_code", "email", "first_name"], "the Campaigns sheet lists the current codes")

# ---------------------------------------------------------------- a lead file arrives
def mail(frm="omar@blacswanken.com", subject="Leads", verified=True, files=()):
    m = EmailMessage()
    m["From"], m["To"], m["Subject"] = frm, "leads@blacswanken.com", subject
    m["Message-ID"] = f"<{len(smtp_out)}-{subject}@x>"
    if verified:
        m["Authentication-Results"] = "mx.google.com; dkim=pass header.d=blacswanken.com; spf=pass; dmarc=pass (p=NONE) header.from=blacswanken.com"
    m.set_content("see attached")
    for name, data in files:
        m.add_attachment(data, maintype="application", subtype="octet-stream", filename=name)
    return m
bot.suppress(con, "gone@unsub.com", "unsubscribe")
con.execute("INSERT INTO leads(email, first_name, status, step, created_at, campaign_id) VALUES ('busy@inseq.com','B','active',1,'2026-01-01',1)")
con.commit()
wb2 = openpyxl.Workbook()
ws = wb2.active
ws.title = "Leads"
ws.append(heads)
rows = [
    ["NEO-SMB-UK", "jane@acme.co.uk", "Jane", "Doe", "Ops Director", "Acme", "", "United Kingdom", "Manchester"],
    ["NEO-US", "joe@corp.us", "Joe", "", "CTO", "Corp", "", "USA", "Seattle"],
    ["NEO-US", "info@corp2.us", "Info", "", "CTO", "Corp2", "", "USA", ""],
    ["NEO-US", "bad-address", "X", "", "CTO", "Corp3", "", "USA", ""],
    ["NEO-US", "joe@corp.us", "Joe", "", "CTO", "Corp", "", "USA", ""],
    ["NOPE", "who@co.com", "W", "", "CEO", "Co", "", "USA", ""],
    ["NEO-US", "missing@co.com", "", "", "", "Co", "", "", ""],
    ["NEO-US", "gone@unsub.com", "G", "", "CEO", "Unsub", "", "USA", ""],
    ["NEO-US", "busy@inseq.com", "B", "", "CEO", "Seq", "", "USA", ""],
    ["NEO-US", "x@nomail.example", "N", "", "CEO", "NoMail", "", "USA", ""],
]
for r in rows:
    ws.append(r)
buf = io.BytesIO()
wb2.save(buf)
print(ops.handle_intake(cfg(), con, mail(frm="someone@else.com", files=[("leads.xlsx", buf.getvalue())])))
check(con.execute("SELECT COUNT(*) FROM leads").fetchone()[0] == 1 and not smtp_out, "a file from anyone but the owner is ignored")
print(ops.handle_intake(cfg(), con, mail(verified=False, files=[("leads.xlsx", buf.getvalue())])))
check(con.execute("SELECT COUNT(*) FROM leads").fetchone()[0] == 1, "a file the mail server couldn't verify (spoofed sender) is refused")
print(ops.handle_intake(cfg(), con, mail(files=[("leads.xlsx", buf.getvalue())])))
got = {r["email"]: r["campaign_id"] for r in con.execute("SELECT email, campaign_id FROM leads")}
check(got.get("jane@acme.co.uk") == 1 and got.get("joe@corp.us") == c2 and len(got) == 3, "good rows go to the campaign named by their code (two campaigns in one file)")
receipt = smtp_out[-1]
body = receipt.get_body(("plain",)).get_content()
att = [p for p in receipt.iter_attachments()]
rej = list(csv.reader(io.StringIO(att[0].get_content() if isinstance(att[0].get_content(), str) else att[0].get_content().decode("utf-8-sig"))))
reasons = " | ".join(r[3] for r in rej[1:])
check(receipt["To"] == "omar@blacswanken.com" and "2 added, 8 rejected" in body, "a receipt goes back with the counts")
check(all(x in reasons for x in ("role address", "bad address format", "duplicate in the file", "unknown campaign code", "missing first_name",
                                 "unsubscribed", "already in an active sequence", "no mail server")),
      f"the rejected-rows file gives the reason for each row ({len(rej) - 1} rows)")
csv_data = "campaign_code,email,first_name,job_title,company,country\nNEO-SMB-UK,kim@kco.com,Kim,CEO,KCo,Ireland\n".encode()
print(ops.handle_intake(cfg(), con, mail(subject="More leads", files=[("more.csv", csv_data)])))
check(con.execute("SELECT 1 FROM leads WHERE email='kim@kco.com'").fetchone(), "CSV files work too")
print(ops.handle_intake(cfg(), con, mail(subject="hello")))
check("no lead file" in smtp_out[-1].get_body(("plain",)).get_content() and list(smtp_out[-1].iter_attachments()), "an email without a file gets the template back")
print(ops.handle_intake(cfg(), con, mail(subject="STOP ALL")))
check(autopilot.killed(con)[0], "STOP ALL from the owner stops all sending")
print(ops.handle_intake(cfg(), con, mail(subject="START ALL now")))
check(not autopilot.killed(con)[0], "START ALL starts it again")
print(ops.handle_intake(cfg(), con, mail(frm="attacker@evil.com", subject="START ALL")))
check(con.execute("SELECT COUNT(*) FROM intake_log").fetchone()[0] >= 7, "every intake email is logged")

# ---------------------------------------------------------------- the 10-day alert
s = bot.load_senders(cfg())[0]
autopilot.rep(con, s["email"])
con.execute("UPDATE inbox_rep SET ready=1, started_at='2026-01-01', full_limit=10")
con.commit()
before = len(smtp_out)
alert = ops.stock_alert(cfg(), con)
m = smtp_out[-1]
check(alert and len(smtp_out) == before + 1 and "Leads running low" in m["Subject"] and "URGENT" in m["Subject"], "below 10 days (here below 3): an urgent email to the owner")
check(any(p.get_filename() == "BSO-lead-template.xlsx" for p in m.iter_attachments()), "the alert has the template attached")
check(ops.stock_alert(cfg(), con) is None, "it isn't repeated within 2 days")
con.execute("UPDATE alerts SET at=?", (bot.iso(bot.now_utc() - timedelta(days=2, hours=2)),))
check(ops.stock_alert(cfg(), con), "it's repeated every 2 days while stock stays low")
for k in range(200):
    con.execute("INSERT INTO leads(email, first_name, status, step, created_at, campaign_id, next_send_at) VALUES (?,?,'active',0,'2026-01-01',?,'2099-01-01')",
                (f"stock{k}@s{k}.com", "S", 1))
    con.execute("INSERT INTO leads(email, first_name, status, step, created_at, campaign_id, next_send_at) VALUES (?,?,'active',0,'2026-01-01',?,'2099-01-01')",
                (f"stockus{k}@u{k}.com", "S", c2))
con.commit()
con.execute("UPDATE alerts SET at='2000-01-01'")
check(ops.stock_alert(cfg(), con) is None, "no alert once stock is above 10 days")

# ---------------------------------------------------------------- daily digest
text = ops.digest(cfg(), con)
check("Sent:" in text and "Lead stock:" in text and "BSO daily digest" in smtp_out[-1]["Subject"], "the daily digest reaches the owner")

# ---------------------------------------------------------------- readiness check of a new inbox
class FakeIMAP:
    def __init__(self, *a, **k): pass
    def login(self, u, p):
        if p != "pw":
            raise Exception("bad login")
    def logout(self): pass
ops.imaplib.IMAP4_SSL = FakeIMAP
import health
health.check = lambda d: {"spf": (True, ""), "dkim": (False, "not found"), "dmarc": (True, "")}
ops.blocklisted = lambda d: ""
sent_tests = []
bot.send_email = lambda *a, **k: sent_tests.append(a) or "<t@x>"
ok, detail = ops.readiness(cfg(), con, s)
check(not ok and "DKIM" in detail and autopilot.limit_for(cfg(), con, s)[0] == 0, "an inbox whose domain fails DKIM is not used")
health.check = lambda d: {"spf": (True, ""), "dkim": (True, ""), "dmarc": (True, "")}
ok, detail = ops.readiness(cfg(), con, s)
check(ok and sent_tests and autopilot.limit_for(cfg(), con, s)[0] > 0, "login, DNS, blocklists and a test send pass: the inbox is ready")

# ---------------------------------------------------------------- answers to information requests
lead_row = con.execute("SELECT * FROM leads WHERE email='jane@acme.co.uk'").fetchone()
con.execute("UPDATE leads SET status='replied', sender=? WHERE id=?", (s["email"], lead_row["id"]))
con.execute("UPDATE campaigns SET info_mode='template', info_text=? WHERE id=1",
            ("Hi {first_name},\n\nHappy to. In short: we provide senior engineers on a monthly basis. Shall I send two short case studies?\n\n{sender_name}",))
con.commit()
replies = []
bot.send_email = lambda s, to, subj, body, in_reply_to=None, dry_run=False, cc=None, headers=None, key=None, lead_id=None, statuses=None, check_dnc=True: replies.append((to, subj, body, in_reply_to)) or "<r@x>"
q = EmailMessage()
q["From"], q["Subject"], q["Message-ID"] = "jane@acme.co.uk", "Re: Development partner for Acme?", "<q1@acme>"
q.set_content("Could you send me some more information about what you do?")
lead = dict(con.execute("SELECT * FROM leads WHERE id=?", (lead_row["id"],)).fetchone())
r = ops.info_reply(cfg(), con, lead, s, q, q.get_content())
check(r and replies and replies[-1][0] == "jane@acme.co.uk" and replies[-1][3] == "<q1@acme>" and "Hi Jane" in replies[-1][2],
      "an information request gets the owner's template answer, in the same thread")
check(con.execute("SELECT COUNT(*) FROM events WHERE type='info_reply'").fetchone()[0] == 1, "the answer is logged")
q2 = EmailMessage()
q2["From"], q2["Subject"] = "jane@acme.co.uk", "Re: x"
q2.set_content("What are your rates, and can we set up a call next week?")
check(ops.info_reply(cfg(), con, lead, s, q2, q2.get_content()) is None, "price and meeting questions still go to a person")
con.execute("UPDATE campaigns SET info_mode='off' WHERE id=1")
check(ops.info_reply(cfg(), con, lead, s, q, q.get_content()) is None, "off by default: nothing is answered")
print(f"\nAll {n} intake checks passed.")
